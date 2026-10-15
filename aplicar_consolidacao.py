# -*- coding: utf-8 -*-
"""PASSO 4 — Consolidação de Horas do cliente, direto do Caderno (aba Consolidação).

Rode:  python aplicar_consolidacao.py
Copia a Consolidação do mês anterior para SAIDA e, em cada aba de PRT:
  - linha 11: datas do ciclo novo (uma por coluna, a partir de H);
  - apaga as horas lançadas no mês passado (só valores digitados; fórmulas ficam);
  - lança as horas do Caderno por dia na linha da função (normal / 75% / 100%).
Duas funções do Caderno que caem na mesma linha (ex.: Soldador Chaparia e MIG -> SOLDADOR)
são SOMADAS.

Criação (tudo listado no relatório), sempre copiando linhas/aba já existentes:
  - PRT sem aba: copia a aba de outra PRT, deixa só um bloco e monta os blocos das funções
    do mês; cabeçalho com o nome da PRT (da Medição);
  - função sem linha: novo bloco (linha da função + linha do total) no fim dos blocos;
  - hora extra 75%/100% sem linha: nova linha logo abaixo da função e o total do bloco soma as duas;
  - blocos sem hora no mês ficam OCULTOS; abas de PRT sem hora no mês ficam OCULTAS;
  - período e data de entrega do cabeçalho (linhas 6-9) atualizados;
  - adicional noturno na linha '<FUNÇÃO> ADICIONAL NOTURNO' (criada se faltar);
  - retroativos (aba Retroativos do Caderno) em colunas de data ANTES do ciclo, com os RDOs nas observações.
Saída (desde o BM 5031): uma planilha por PRT, só com a aba dela, na pasta da PRT; a de várias abas vai para
_arquivo\\consolidacao_varias_abas (é o modelo do mês seguinte).
Não vão para a Consolidação: montador de andaime / operador de PTA (sem item próprio) e
mobilização, manutenção e fornecimento (são verba/volume).
"""
import glob
import os
import re
import shutil
import sys
from datetime import date, datetime, time, timedelta

import openpyxl
from openpyxl.utils.datetime import to_excel

import config_mes as C
from comum import (fonte, HE_DO_OPERADOR, HE_TIPOS, SEM_ITEM, Excel, aba_acumulado, padronizar_fontes, arquivo,
                   data_excel, he_vira_normal, norm, pasta_prt, prt_info)

# nome no Caderno (sem acento) -> inícios de rótulo aceitos na aba do cliente, em ordem de preferência.
# O próprio nome do Caderno é sempre tentado primeiro.
APELIDOS = {
    "TECNICO EM SEGURANCA DO TRABALHO": ["TEC SEGURANCA", "TECNICO SEGURANCA", "TECNICO DE SEGURANCA"],
    "TECNICO DE PLANEJAMENTO": ["TEC PLANEJAMENTO"],
    "TECNICO DE QUALIDADE": ["TEC QUALIDADE"],
    "AUXILIAR ADMINISTRATIVO": ["AUXILIAR ADM"],
    "SUPERVISOR DE ELETRICA": ["SUPERVISOR ELETROMECANICO", "SUPERVISOR ELETRICA"],
    "CALDEIREIRO": ["CALDEREIRO"],
    "ELETRICISTA FC": ["ELETRICISTA", "ELETRCISTA"],
    "MECANICO MONTADOR": ["MONTADOR INDUSTRIAL"],
    "SOLDADOR CHAPARIA": ["SOLDADOR"],
    "SOLDADOR MIG": ["SOLDADOR"],
    "ENCARREGADO DE ANDAIMES": ["ENCARREGADO DE ANDAIME"],
    "ENCARREGADO DE MECANICA": ["ENCARREGADO DE MEC"],
    "ENCARREGADO DE ELETRICA": ["ENCARREGADO DE ELE"],
    "INSPETOR DE SOLDA N1/VISUAL E LIQUIDIO PENETRANTE": ["INSPETOR DE SOLDA"],
    "OPERADOR DE CAMINHAO MUCK": [],          # a hora aparece na linha do caminhão
    # equipamentos (quantidade por dia)
    "ANDAIME": ["ANDAIME"],
    "GUINDASTE 80 TON COM OPERADOR E SINALEIRO": ["GUINDASTE 80"],
    "MICRO ONIBUS 28 LUGARES INCLUIDO O MOTORISTA": ["MICRO ONIBUS"],
    "PLATAFORMA ELEVATORIA 25 METROS OPERADA POR ELETRICISTA/MECANICO": ["PLATAFORMA 25", "PLAT ELEVATORIA PROD 25"],
    "PLATAFORMA ELEVATORIA 15 METROS OPERADA POR ELETRICISTA/MECANICO": ["PLATAFORMA 15", "PLAT ELEVATORIA PROD 15"],
    "CAMINHONETE 4X4": ["CAMINHONETE"],
    "CAMINHAO GUINDAUTO 10 TON INCLUINDO OPERADOR COM JOYSTICK REMOTO": ["CAMINHAO MUNCK", "CAMINHAO GUINDAUTO"],
    "VAN 15 LUGARES": ["VAN 15"],
}
TIPO = {"Normais": "N", "H. Extras 75%": "75", "H. Extras 100%": "100", "Ad. Noturno": "NOT",
        "Equipamento": "N", "Guindaste": "N", "Veículo": "N", "Plataforma": "N", "Munck": "N"}
ROTULO_TIPO = {"75": "HORAS 75 %", "100": "HORAS 100%", "NOT": "ADICIONAL NOTURNO"}
TIPOS_HORA = {"Normais", "H. Extras 75%", "H. Extras 100%", "Ad. Noturno"}
TIPO_TXT = {"Normais": "horas normais", "H. Extras 75%": "horas extras", "H. Extras 100%": "horas extras",
            "Ad. Noturno": "adicional noturno"}
ARQ = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_arquivo", "consolidacao_varias_abas")
XL_RIGHT = -4161
XL_SEM_COR = -4142
SEM_CONSOLIDACAO = {"Volume"}   # mobilização, manutenção, fornecimento: não vão para a Consolidação
FORA_DA_CONSOLIDACAO = ("MANUTENCAO", "MOBILIZACAO", "FORNECIMENTO")
# rótulo de linha nova (equipamentos têm nome curto na Consolidação)
ROTULO_NOVO = {
    "CAMINHAO GUINDAUTO 10 TON INCLUINDO OPERADOR COM JOYSTICK REMOTO": "CAMINHÃO MUNCK",
    "MICRO ONIBUS 28 LUGARES INCLUIDO O MOTORISTA": "MICRO-ÔNIBUS 28 LUGARES",
    "GUINDASTE 80 TON COM OPERADOR E SINALEIRO": "GUINDASTE 80 TON",
}
EQUIP = ("ONIBUS", "MUNCK", "PLAT", "GUINDASTE", "CAMINHONETE", "VAN ", "ANDAIME", "GERADOR", "CAMINHAO")
XL_DOWN = -4121


def achar_grupo(grupos, func):
    # linhas de adicional noturno nunca recebem hora normal/extra (ex.: 'CALDEIREIRO ADICIONAL
    # NOTURNO' tem a grafia certa e a linha normal está 'CALDEREIRO')
    grupos = {k: v for k, v in grupos.items() if "ADICIONAL" not in k and "NOTURN" not in k}
    n = norm(func)
    for cand in [n] + APELIDOS.get(n, []):
        if cand in grupos:
            return cand
        pref = sorted((k for k in grupos if k.startswith(cand)), key=len)
        if pref:
            return pref[0]
    return None
COL_DIA1 = 8  # H


def horas(v):
    """Valor da célula em horas (ou quantidade, se for número puro). Células formatadas
    como hora chegam como time; acima de 24h o openpyxl devolve datetime de 1900 —
    converte pelo serial do Excel."""
    if isinstance(v, bool) or v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, time):
        return v.hour + v.minute / 60 + v.second / 3600
    if isinstance(v, timedelta):
        return v.total_seconds() / 3600
    if isinstance(v, datetime):
        return to_excel(v) * 24
    return 0.0


def ler_caderno_consolidacao(chave_bloco=None):
    """{prt: {(funcao, tipo): {date: horas}}}, avisos. chave_bloco(título do bloco) troca o nº da PRT por outra chave
    (consolidacao_preparacao: a equipe); bloco com chave None fica de fora."""
    ws = openpyxl.load_workbook(C.CADERNO, data_only=True)["Consolidação"]
    dados, avisos, prt, cols = {}, [], None, {}
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a and norm(a) == "FUNCAO":
            cols = {}
            for c in range(3, ws.max_column + 1):
                v = str(ws.cell(r, c).value or "")
                m = re.fullmatch(r"(\d\d)/(\d\d)", v.strip())
                if m:
                    d, mth = int(m.group(1)), int(m.group(2))
                    ano = C.PERIODO_FIM.year if mth <= C.PERIODO_FIM.month else C.PERIODO_INI.year
                    cols[c] = date(ano, mth, d)
            continue
        if a and ws.cell(r, 2).value is None and "CLIENTE" in str(a).upper():
            m = re.search(r"PRT_?\s*(\d{4})", str(a))
            prt = chave_bloco(str(a)) if chave_bloco else (m.group(1) if m else None)
            if not prt and not chave_bloco:
                avisos.append(f"bloco '{a}' não é de PRT — não vai para a Consolidação do cliente")
            continue
        if prt and a and ws.cell(r, 2).value:
            chave = (str(a).strip(), str(ws.cell(r, 2).value).strip())
            por_dia = {d: horas(ws.cell(r, c).value) for c, d in cols.items() if horas(ws.cell(r, c).value)}
            if por_dia:
                dados.setdefault(prt, {})[chave] = por_dia
    return dados, avisos


def com_retroativos(dados, retro):
    """Horas do ciclo + retroativos, já com as regras de cobrança (as mesmas da Base):
    - hora extra do operador de Munck vai para a linha do caminhão como hora (Munck não tem hora extra);
    - hora extra de função sem hora extra na proposta vira hora normal (HE_COMO_NORMAL: 5042, 5029, 5007 Inspetor)."""
    avisos = []
    for prt, por in retro.items():
        for chave, por_dia in por.items():
            alvo = dados.setdefault(prt, {}).setdefault(chave, {})
            for d, h in por_dia.items():
                alvo[d] = alvo.get(d, 0) + h
    for prt, por in dados.items():
        for (func, tipo), por_dia in list(por.items()):
            alvo = HE_DO_OPERADOR.get(norm(func))
            if alvo and tipo.replace(" (pendência)", "") in HE_TIPOS:
                cam = next((k for k in por if norm(k[0]).startswith(alvo)), None)
                if cam is None:
                    avisos.append(f"PRT {prt} {func} / {tipo}: hora extra do operador sem a linha do equipamento")
                    continue
                for d, h in por_dia.items():
                    por[cam][d] = por[cam].get(d, 0) + h
            elif he_vira_normal(prt, func, tipo):
                normal = por.setdefault((func, "Normais"), {})
                for d, h in por.pop((func, tipo)).items():
                    normal[d] = normal.get(d, 0) + h
    return dados, avisos


def grupos_da_aba(ws_py):
    """{rotulo_base_normalizado: {"N"/"75"/"100": linha}} a partir da coluna B."""
    grupos, atual = {}, None
    for r in range(13, min(ws_py.max_row, 200) + 1):
        b = ws_py.cell(r, 2).value
        if not b:
            continue
        n = norm(b)
        if "100" in n or re.search(r"HORAS? 10\b", n):
            tipo = "100"
        elif "75" in n:
            tipo = "75"
        else:
            tipo = "N"
        if tipo == "N":
            atual = n
            grupos[atual] = {"N": r}
        elif atual:
            grupos[atual].setdefault(tipo, r)
    return grupos


def aba_da_prt(wb_py, prt):
    cands = [ws for ws in wb_py.worksheets if prt in ws.title.replace(" ", "")]
    vis = [ws for ws in cands if ws.sheet_state == "visible"]
    return (vis or cands or [None])[0]


# ----------------------------------------------------------------- estrutura e criação (COM)
def tipo_rotulo(n):
    if "NOTURN" in n:           # 'CALDEIREIRO ADICIONAL NOTURNO' é linha do bloco do caldeireiro ('HORAS NOTURNAS' também)
        return "NOT"
    if "100" in n or re.search(r"HORAS? 10\b", n):
        return "100"
    # 1ª faixa de HE: '75' nas PRTs; 'HORAS 50 %' na Consolidação da Manutenção (contrato 4600000002)
    return "75" if "75" in n or re.search(r"HORAS? 50\b", n) else "N"


def col(c):
    s = ""
    while c:
        c, r = divmod(c - 1, 26)
        s = chr(65 + r) + s
    return s


def estrutura(ws):
    """Lê a aba pelo Excel -> (coluna TOTAL, coluna TOTAL GERAL, linha DISCRIMINAÇÃO,
    blocos [{rot, linhas{N/75/100: linha}, ini, ult, fim}])."""
    r11 = ws.Range("A11:CZ11").Value[0]
    tots = [i + 1 for i, v in enumerate(r11) if norm(v or "").startswith("TOTAL")]
    soma, geral = tots[0], (tots[1] if len(tots) > 1 else tots[0] + 1)
    colb = [v for (v,) in ws.Range("B1:B300").Value]
    disc = next(i + 1 for i, v in enumerate(colb) if i >= 12 and norm(v or "").startswith("DISCRIMINA"))
    blocos, atual = [], None
    for r in range(13, disc):
        b = colb[r - 1]
        if not b:
            continue
        n = norm(b)
        t = tipo_rotulo(n)
        if t == "N" or atual is None:
            atual = {"rot": n, "linhas": {t: r}, "ini": r}
            blocos.append(atual)
        else:
            atual["linhas"].setdefault(t, r)
        atual["ult"] = r
    for i, bl in enumerate(blocos):
        prox = blocos[i + 1]["ini"] if i + 1 < len(blocos) else disc
        fim = bl["ult"]
        if fim + 1 < prox and not colb[fim] and ws.Cells(fim + 1, geral).HasFormula:
            fim += 1
        bl["fim"] = fim
        bl["oculta"] = bool(ws.Rows(bl["ini"]).Hidden)
    return soma, geral, disc, blocos


def por_rotulo(blocos):
    """{rótulo: bloco}. Rótulo repetido (5008: 'MICRO ONIBUS 28 LUGARES' antiga e oculta, 'MICRO-ÔNIBUS 28 LUGARES' em
    uso com o total de diárias) -> vale o bloco VISÍVEL no modelo; senão o primeiro."""
    out = {}
    for bl in blocos:
        atual = out.get(bl["rot"])
        if atual is None or (atual.get("oculta") and not bl.get("oculta")):
            out[bl["rot"]] = bl
    return out


def limpar_dias(ws, r, soma):
    rng = ws.Range(ws.Cells(r, COL_DIA1), ws.Cells(r, soma - 1))
    f = rng.Formula[0]
    if any(x not in (None, "") and not str(x).startswith("=") for x in f):
        rng.Formula = (tuple(x if str(x).startswith("=") else "" for x in f),)


def copiar_linha(xl, ws, src, dest):
    ws.Rows(src).Copy()
    ws.Rows(dest).Insert(XL_DOWN)   # argumentos POSICIONAIS: por nome o pywin32 ignora
    xl.app.CutCopyMode = False
    ws.Rows(dest).Hidden = False


def novo_bloco(xl, ws, rotulo, equip):
    """Linha da função + linha do total, copiadas de um bloco simples da própria aba."""
    soma, geral, _, blocos = estrutura(ws)
    simples = [b for b in blocos if len(b["linhas"]) == 1]
    pref = [b for b in simples if any(e in b["rot"] + " " for e in EQUIP) == equip]
    mod = (pref or simples or blocos)[0]
    pos = max(b["fim"] for b in blocos) + 1
    copiar_linha(xl, ws, mod["ini"], pos)
    if mod["fim"] > mod["ult"]:
        copiar_linha(xl, ws, mod["fim"], pos + 1)
        f_mod = str(ws.Cells(pos + 1, geral).Formula)
        ws.Cells(pos + 1, geral).Formula = (f"={col(soma)}{pos}/30" if "/30" in f_mod
                                           else f"=SUM({col(soma)}{pos}:{col(soma)}{pos})*24")
        rot_tot = ws.Cells(pos + 1, soma).Value       # 'TOTAL DE HORAS DE <função do bloco copiado>' (5029: Soldador TIG
        if isinstance(rot_tot, str) and norm(rot_tot).startswith("TOTAL DE HORAS"):   # saía como Supervisor de Elétrica)
            ws.Cells(pos + 1, soma).Value = f"TOTAL DE HORAS DE {rotulo}"
    ws.Cells(pos, 2).Value = rotulo
    limpar_dias(ws, pos, soma)
    return pos


def nova_extra(xl, ws, bl, t, soma, geral):
    """Linha de hora extra logo abaixo das linhas do bloco; o total do bloco soma todas."""
    pos = bl["ult"] + 1
    copiar_linha(xl, ws, bl["linhas"].get("N", bl["ini"]), pos)
    base = str(ws.Cells(bl["ini"], 2).Value or "").strip()
    ws.Cells(pos, 2).Value = f"{base} {ROTULO_TIPO[t]}"
    limpar_dias(ws, pos, soma)
    if bl["fim"] > bl["ult"]:
        tot = bl["fim"] + 1
        ws.Cells(tot, geral).Formula = f"=SUM({col(soma)}{bl['ini']}:{col(soma)}{pos})*24"
    return pos


def num_aba(ws):
    m = re.search(r"PRT\s*_?(\d{4})", ws.Name)
    return int(m.group(1)) if m else None


def aba_com(wb, prt):
    cands = [ws for ws in wb.Worksheets if prt in ws.Name.replace(" ", "") and num_aba(ws)]
    vis = [ws for ws in cands if ws.Visible == -1]
    return (vis or cands or [None])[0]


def nova_aba_prt(xl, wb, prt, lg):
    """Copia a aba visível de outra PRT (a de menos blocos), deixa um bloco só e acerta o cabeçalho."""
    vis = [ws for ws in wb.Worksheets if ws.Visible == -1 and num_aba(ws)]
    modelo = min(vis, key=lambda w: len(estrutura(w)[3]))
    antes = [w for w in vis if num_aba(w) < int(prt)]
    nomes = {s.Name for s in wb.Sheets}
    if antes:
        modelo.Copy(None, max(antes, key=num_aba))   # (Before, After) posicionais
    else:
        modelo.Copy(min(vis, key=lambda w: w.Index))
    nova = next(s for s in wb.Sheets if s.Name not in nomes)   # a cópia é a única aba com nome novo
    nova.Name = f"PRT {prt}"
    med = fonte(prt, "medicao")
    ident, nome = prt_info(med) if med else (f"PRT_{prt}", "")
    titulo = f"{ident} - {nome}"
    acu = fonte(prt, "acumulado")
    if acu:   # título oficial da PRT (A6 do Acumulado); a Medição do mês passado pode vir sem o nome (5001)
        wa = aba_acumulado(openpyxl.load_workbook(acu, data_only=True))
        titulo = str(wa["A6"].value or "").strip() or titulo
    velho = str(num_aba(modelo))
    cab = nova.Range("A1:CZ10").Value
    for i, row in enumerate(cab):
        for j, v in enumerate(row):
            if isinstance(v, str) and velho in v and not nova.Cells(i + 1, j + 1).HasFormula:
                novo = titulo if norm(v).startswith("PRT") else v.replace(velho, prt)   # como nas outras abas
                nova.Cells(i + 1, j + 1).Value = novo
    soma, geral, disc, blocos = estrutura(nova)
    simples = [b for b in blocos if len(b["linhas"]) == 1 and not any(e in b["rot"] + " " for e in EQUIP)]
    fica = (simples or blocos)[0]
    manter = set(range(fica["ini"], fica["fim"] + 1))
    for r in range(max(b["fim"] for b in blocos), 12, -1):
        if r not in manter:
            nova.Rows(r).Delete()
    nova.Rows(f"13:{12 + len(manter)}").Hidden = False
    _, _, disc, _ = estrutura(nova)
    for r in range(disc + 1, disc + 30):
        v = nova.Cells(r, 3).Value
        if isinstance(v, str) and "APONTAMENTO CONFORME RDO" not in norm(v):
            nova.Cells(r, 3).Value = None
    lg.append(f"CRIADA aba '{nova.Name}' (cópia de '{modelo.Name}'), cabeçalho: {titulo}")
    return nova


def ler_retroativos():
    """Aba Retroativos do Caderno -> ({prt: {(funcao, tipo): {date: horas}}}, {prt: [(funcao, tipo, rdo)]}, avisos)."""
    wb = openpyxl.load_workbook(C.CADERNO, data_only=True)
    if "Retroativos" not in wb.sheetnames:
        return {}, {}, []
    dados, itens, avisos = {}, {}, []
    for prt, func, tipo, d, h, rdo in wb["Retroativos"].iter_rows(min_row=5, max_col=6, values_only=True):
        if not prt or not func:
            continue
        prt, func, tipo = str(prt).strip(), str(func).strip(), str(tipo).strip()
        if not isinstance(d, datetime):
            avisos.append(f"PRT {prt} {func} / {tipo} = {h}: retroativo sem os dias no Caderno — não entrou")
            continue
        por = dados.setdefault(prt, {}).setdefault((func, tipo), {})
        por[d.date()] = por.get(d.date(), 0) + float(h)
        itens.setdefault(prt, []).append((func, tipo, str(rdo or "").strip()))
    return dados, itens, avisos


def ajustar_colunas(xl, ws, n, lg):
    """A aba fica com n colunas de dia a partir de H (retroativos + ciclo). Sobrando (retroativos do mês passado),
    saem colunas SEM texto fora da grade (cabeçalho/assinaturas ficam); faltando, entram cópias de uma coluna de
    dia DENTRO do intervalo, e os SUM(H:AL) das linhas se esticam."""
    soma, _, disc, _ = estrutura(ws)
    tem = soma - COL_DIA1
    if tem > n:
        fora = [r for r in list(range(1, 11)) + list(range(disc, disc + 40))]
        livres = [c for c in range(COL_DIA1 + 1, soma) if not any(ws.Cells(r, c).Value not in (None, "") for r in fora)]
        if len(livres) < tem - n:
            lg.append(f"!! só {len(livres)} coluna(s) livre(s) para tirar {tem - n} — dias a mais ficam sem data")
        for c in sorted(livres[:tem - n], reverse=True):
            ws.Columns(c).Delete()
        lg.append(f"{min(len(livres), tem - n)} coluna(s) de dia a mais tirada(s) (retroativos do mês passado)")
    elif tem < n:
        # cópia de uma coluna sem texto no cabeçalho (copiar a I levava o I2 '=PRT5009…' junto: '0 0 0…' no topo da 5007)
        fora = list(range(1, 11)) + list(range(disc, disc + 40))
        livres = [c for c in range(COL_DIA1 + 1, soma) if not any(ws.Cells(r, c).Value not in (None, "") for r in fora)]
        alvo = livres[0] if livres else COL_DIA1 + 1
        for _ in range(n - tem):
            ws.Columns(alvo).Copy()
            ws.Columns(alvo).Insert(XL_RIGHT)   # posicional: por nome o pywin32 ignora
            xl.app.CutCopyMode = False
        lg.append(f"{n - tem} coluna(s) de dia criada(s) para os retroativos")


def referencias(ws):
    """ANTES de trocar as datas: largura e formato de data das colunas de dia do modelo e o cinza dos fins de semana."""
    from collections import Counter
    soma, _, _, blocos = estrutura(ws)
    cols = list(range(COL_DIA1, soma))
    datas = {c: date(1899, 12, 30) + timedelta(days=int(v)) for c in cols
             for v in [ws.Cells(11, c).Value2] if isinstance(v, float) and v > 0}
    linhas = [r for bl in blocos for r in bl["linhas"].values()][:8]
    cinza = Counter(ws.Cells(r, c).Interior.Color for c, d in datas.items() if d.weekday() >= 5 for r in linhas
                    if ws.Cells(r, c).Interior.Pattern != XL_SEM_COR)
    return {"largura": Counter(round(ws.Columns(c).ColumnWidth, 2) for c in cols).most_common(1)[0][0],
            "fmt": (Counter(str(ws.Cells(11, c).NumberFormat) for c in datas).most_common(1) or [("dd/mm", 0)])[0][0],
            "cinza": cinza.most_common(1)[0][0] if cinza else 14277081}       # D9D9D9


def formatar_dias(ws, ref, dias_tab):
    """Colunas de dia com a cara do modelo — mesma largura, data no mesmo formato — e cinza só em sábado/domingo
    das linhas de função: o cinza do modelo era fixo na coluna e, com as datas novas, caía em dia útil (06/10/2026)."""
    _, _, disc, blocos = estrutura(ws)
    linhas = sorted({r for bl in blocos for r in bl["linhas"].values()})
    for i, d in enumerate(dias_tab):
        c = COL_DIA1 + i
        ws.Columns(c).ColumnWidth = ref["largura"]
        ws.Cells(11, c).NumberFormat = ref["fmt"]
        ws.Range(ws.Cells(13, c), ws.Cells(disc - 1, c)).Interior.Pattern = XL_SEM_COR
        if d.weekday() >= 5:
            for r in linhas:
                ws.Cells(r, c).Interior.Color = ref["cinza"]


def faixa_da_linha(f, r, ini, fim):
    """Fórmula do total da linha r com o intervalo de dias trocado por ini:fim (só intervalos da própria linha)."""
    return re.sub(rf"\$?[A-Z]{{1,3}}\$?{r}:\$?[A-Z]{{1,3}}\$?{r}\b", f"{ini}{r}:{fim}{r}", f)


def totais_da_linha(ws):
    """Total de cada linha (coluna TOTAL) soma TODAS as colunas de dia — no BM030 as linhas com retroativo
    somavam desde H e as outras só o ciclo (V:AZ); somar tudo dá o mesmo e não esquece retroativo."""
    soma, _, disc, _ = estrutura(ws)
    ini, fim = col(COL_DIA1), col(soma - 1)
    for r in range(13, disc):
        f = str(ws.Cells(r, soma).Formula)
        if f.startswith("="):
            novo = faixa_da_linha(f, r, ini, fim)
            if novo != f:
                ws.Cells(r, soma).Formula = novo


def e_lista(xs):
    xs = list(dict.fromkeys(xs))
    return xs[0] if len(xs) == 1 else ", ".join(xs[:-1]) + " e " + xs[-1]


def texto_retroativo(itens):
    """[(função, tipo, rdo)] -> frase das observações, no estilo do BM030 ('… RDOs 5313; 5322')."""
    por_func = {}
    for f, t, _ in itens:
        por_func.setdefault(f, []).append(TIPO_TXT.get(t, t))
    rdos = "; ".join(sorted({r for _, _, r in itens if r}, key=int))
    if all(set(ts) == {"horas normais"} for ts in por_func.values()):
        um = len(por_func) == 1
        txt = (f"Apontamentos retroativos da{'' if um else 's'} {'função' if um else 'funções'} de {e_lista(sorted(por_func))}, "
               f"que não {'estava prevista' if um else 'estavam previstas'} na PRT.")
    else:
        partes = "; ".join(f"{f} ({e_lista(ts)})" for f, ts in sorted(por_func.items()))
        txt = f"Apontamentos retroativos que não estavam previstos na PRT: {partes}."
    return f"{txt} RDOs {rdos}" if rdos else txt


def observacoes(ws, itens, lg):
    """Texto de retroativo do mês passado sai, e também a observação de equipamento sem uso no mês (5029: 'VAN 15
    LUGARES - VERBA MENSAL' sem van medida); o retroativo deste mês (funções + RDOs) entra depois das outras."""
    soma, _, disc, blocos = estrutura(ws)
    usa = lambda bl: any(any(v not in (None, "", 0) for v in ws.Range(ws.Cells(r, COL_DIA1), ws.Cells(r, soma - 1)).Value[0])
                         for r in bl["linhas"].values())
    usados = {bl["rot"] for bl in blocos if usa(bl)}      # rótulo repetido (5008: 2 micro-ônibus) vale se um deles foi usado
    parados = {bl["rot"] for bl in blocos if any(e in bl["rot"] + " " for e in EQUIP)} - usados
    col_c = [v for (v,) in ws.Range(f"C{disc + 1}:C{disc + 14}").Value]
    for i, v in enumerate(col_c):
        n = norm(v) if isinstance(v, str) else ""
        if n.startswith("APONTAMENTOS RETROATIVOS") or any(p in n for p in parados if len(p) > 5):
            ws.Cells(disc + 1 + i, 3).Value = None
            col_c[i] = None
            lg.append(f"observação apagada (retroativo do mês passado / equipamento sem uso no mês): {v[:60]}…")
    if itens:
        usadas = [i for i, v in enumerate(col_c) if v not in (None, "")]
        r = disc + 1 + (usadas[-1] + 2 if usadas else 1)
        ws.Cells(r, 3).Value = texto_retroativo(itens)
        lg.append(f"observação C{r}: {texto_retroativo(itens)}")


def data_entrega(ws):
    """Data de entrega do cabeçalho (célula abaixo do rótulo) = data de postagem do BM (config_mes.DATA_BM)."""
    if not getattr(C, "DATA_BM", None):
        return
    for i, row in enumerate(ws.Range("A6:CZ9").Value):
        for j, v in enumerate(row):
            if isinstance(v, str) and norm(v).startswith("DATA DE ENTREGA"):
                ws.Cells(7 + i, j + 1).Value = data_excel(C.DATA_BM)
                return


def exportar(xl, wb, prt, lg):
    """Aba da PRT vira uma planilha só dela, na pasta da PRT em SAIDA (aberta no Excel do usuário: não regrava)."""
    ini, fim = C.PERIODO_INI.strftime("%d-%m"), C.PERIODO_FIM.strftime("%d-%m-%Y")
    dst = os.path.join(pasta_prt(C.SAIDA, prt, C.PASTA_DA_PRT),
                       f"BM GSV- CONSOLIDAÇÃO DE HORAS - {ini} a {fim} - BM{C.BM_NUMERO[-3:]} - PRT {prt}.xlsx")
    if exportar_aba(xl, aba_com(wb, prt), dst):
        lg.append(f"planilha da PRT: {os.path.relpath(dst, C.SAIDA)}")
    else:
        lg.append(f"!! {os.path.relpath(dst, C.SAIDA)}: aberta no Excel — NÃO regravei (feche e rode de novo)")


def exportar_aba(xl, ws, dst):
    """A aba vira uma pasta de trabalho só dela em dst. False = dst aberta no Excel do usuário (não regrava)."""
    if os.path.exists(os.path.join(os.path.dirname(dst), "~$" + os.path.basename(dst))):
        return False
    ur = ws.UsedRange
    for i, linha in enumerate(ur.Formula):
        for j, x in enumerate(linha):
            if isinstance(x, str) and x.startswith("=") and "!" in x:     # lê outra aba (='PRT5009-INDIRETO'!H2)
                cel = ws.Cells(ur.Row + i, ur.Column + j)
                v = cel.Value2                                            # Value2: data fica número (sem 03:00)
                cel.Value2 = None if v in (0, "") else v                  # 0 de célula vazia apareceria no topo
    ws.Copy()                                 # sem argumentos: pasta de trabalho nova só com esta aba
    nb = xl.app.ActiveWorkbook
    for n in list(nb.Names):                  # nomes velhos que vieram com a aba (#REF!, outro arquivo)
        try:
            if "[" in str(n.RefersTo) or "#REF" in str(n.RefersTo):
                n.Delete()
        except Exception:
            pass
    for l in nb.LinkSources(1) or ():
        nb.BreakLink(l, 1)
    nb.SaveAs(dst, 51)
    nb.Close(False)
    return True


def main():
    so = [a for a in sys.argv[1:] if re.fullmatch(r"\d{4}", a)]   # grava só a planilha destas PRTs (as outras ficam)
    dados, avisos = ler_caderno_consolidacao()
    retro, retro_itens, avisos_r = ler_retroativos()
    dados, avisos_c = com_retroativos(dados, retro)        # retroativo entra junto com as horas do ciclo + regras
    avisos += avisos_r + avisos_c
    fonte = ([C.MODELO_CONSOLIDACAO] if getattr(C, "MODELO_CONSOLIDACAO", None) else
             [f for f in glob.glob(os.path.join(C.FONTE_ACUMULADO, "*CONSOLIDA*.xls*")) if "MANUTEN" not in norm(os.path.basename(f))])
    if not fonte:   # desde o BM 5031 sai uma por PRT: o modelo de várias abas do mês passado está no _arquivo
        m = re.search(r"BM (\d{4})", C.FONTE_ACUMULADO)
        fonte = glob.glob(os.path.join(ARQ, f"*CONSOLIDA*BM{m.group(1)[-3:]}.xlsx")) if m else []
    if len(fonte) != 1:
        raise SystemExit(f"Esperava 1 Consolidação em {C.FONTE_ACUMULADO}, achei: {fonte}")
    ini, fim = C.PERIODO_INI.strftime("%d-%m"), C.PERIODO_FIM.strftime("%d-%m-%Y")
    os.makedirs(C.SAIDA, exist_ok=True)
    dst = os.path.join(C.SAIDA, f"BM GSV- CONSOLIDAÇÃO DE HORAS - {ini} a {fim} - BM{C.BM_NUMERO[-3:]}.xlsx")
    shutil.copy2(fonte[0], dst)

    dias = [C.PERIODO_INI + timedelta(d) for d in range((C.PERIODO_FIM - C.PERIODO_INI).days + 1)]
    periodo = f"{C.PERIODO_INI:%d/%m/%Y} A {C.PERIODO_FIM:%d/%m/%Y}"
    log = [f"Modelo: {fonte[0]}", *[f"  ! {a}" for a in avisos]]
    feitas = []
    with Excel() as xl:
        wb = xl.abrir(dst)
        xl.app.ScreenUpdating = False
        for prt in sorted(dados):
            print(f"PRT {prt}...", flush=True)
            lg = []
            if not any(TIPO.get(t) and t not in SEM_CONSOLIDACAO and norm(f) not in SEM_ITEM
                       and not norm(f).startswith(FORA_DA_CONSOLIDACAO) for f, t in dados[prt]):
                log.append(f"\nPRT {prt}: nada que vá para a Consolidação (só volume/fornecimento) — sem aba")
                continue
            ws = aba_com(wb, prt)
            vazia = aba_nova = ws is None
            if aba_nova:
                ws = nova_aba_prt(xl, wb, prt, lg)
            ws.Visible = -1
            # colunas de dia: retroativos (antes do ciclo) + ciclo
            dias_tab = sorted({d for por in retro.get(prt, {}).values() for d in por if d < C.PERIODO_INI}) + dias
            ajustar_colunas(xl, ws, len(dias_tab), lg)
            ref = referencias(ws)                 # antes de trocar as datas (o cinza do modelo é o das datas velhas)
            soma_c, geral_c, _, blocos = estrutura(ws)
            colmap = {d: COL_DIA1 + i for i, d in enumerate(dias_tab)}
            for i, c in enumerate(range(COL_DIA1, soma_c)):
                cel = ws.Cells(11, c)
                if not cel.HasFormula:
                    cel.Value = data_excel(dias_tab[i]) if i < len(dias_tab) else None   # datetime pelo COM grava 03:00
            formatar_dias(ws, ref, dias_tab)
            cab = ws.Range("A6:CZ9").Value
            for i, row in enumerate(cab):
                for j, v in enumerate(row):
                    if isinstance(v, str) and re.search(r"\d\d/\d\d/\d{4} A \d\d/\d\d/\d{4}", v):
                        ws.Cells(6 + i, j + 1).Value = re.sub(r"\d\d/\d\d/\d{4} A \d\d/\d\d/\d{4}", periodo, v)
            data_entrega(ws)
            for bl in blocos:
                for r in bl["linhas"].values():
                    limpar_dias(ws, r, soma_c)
            tem_extra = any(len(bl["linhas"]) > 1 for bl in blocos)

            # 1) garante a estrutura (cria blocos/linhas que faltam)
            pedidos, faltas, e_hora = {}, [], {}
            for (func, tipo), por_dia in dados[prt].items():
                tot = f"{sum(por_dia.values()):g}"
                n = norm(func)
                if (tipo in SEM_CONSOLIDACAO or (n in APELIDOS and not APELIDOS[n]) or n in SEM_ITEM
                        or n.startswith(FORA_DA_CONSOLIDACAO)):
                    continue
                t = TIPO.get(tipo)
                if t is None:
                    faltas.append(f"{func} / {tipo}: tipo não existe na Consolidação do cliente ({tot})")
                    continue
                grupos = por_rotulo(blocos)
                g = achar_grupo(grupos, func)
                if g is None:
                    rotulo = ROTULO_NOVO.get(n) or str(func).split("(")[0].strip().upper()
                    if vazia:
                        bl = blocos[0]
                        ws.Cells(bl["ini"], 2).Value = rotulo
                        vazia = False
                    else:
                        novo_bloco(xl, ws, rotulo, any(e in norm(rotulo) + " " for e in EQUIP))
                    lg.append(f"CRIADO bloco '{rotulo}' ({tot} em {tipo})")
                    soma_c, geral_c, _, blocos = estrutura(ws)
                    g = achar_grupo(por_rotulo(blocos), func)
                bl = por_rotulo(blocos)[g]
                vazia = False   # o bloco que sobrou da cópia já está em uso
                if t not in bl["linhas"]:
                    if tem_extra or aba_nova or t == "NOT":    # noturno nunca vai para a linha normal
                        nova_extra(xl, ws, bl, t, soma_c, geral_c)
                        lg.append(f"CRIADA linha '{g} {ROTULO_TIPO[t]}' ({tot}h) e o total do bloco soma as linhas")
                        soma_c, geral_c, _, blocos = estrutura(ws)
                    else:
                        t = "N"
                        faltas.append(f"{func} / {tipo}: aba não usa linhas de hora extra — somei na linha normal '{g}' ({tot}h)")
                for d, h in por_dia.items():
                    pedidos[(g, t, d)] = pedidos.get((g, t, d), 0) + h
                e_hora[(g, t)] = e_hora.get((g, t), False) or tipo in TIPOS_HORA

            # 2) lança os valores nas linhas finais
            por_rot = por_rotulo(blocos)
            for (g, t, d), h in pedidos.items():
                if d not in colmap:
                    faltas.append(f"{g} / {t}: dia {d:%d/%m/%Y} fora do ciclo e sem retroativo — não lançado ({h:g})")
                    continue
                r = por_rot[g]["linhas"][t]
                cel = ws.Cells(r, colmap[d])
                # hora ou quantidade: hora de função (normal/extra/noturno) é sempre hora; equipamento segue o formato
                # do TOTAL da linha. A coluna pode ter vindo de retroativo do mês passado com outro formato (5029:
                # 'Geral' -> 9 virava 9 dias no total ×24; 5042: 'h:mm' mostrava 27h como 03:00) — hora fica '[h]:mm'
                ftot = str(ws.Cells(r, soma_c).NumberFormat).lower() + str(cel.NumberFormat).lower()
                if e_hora.get((g, t)) or ("h" in ftot and ":" in ftot):
                    cel.Value = h / 24
                    if str(cel.NumberFormat) != "[h]:mm":     # '[h]:mm:ss' dava 18:00:00 e '#####' na coluna estreita
                        cel.NumberFormat = "[h]:mm"
                else:
                    cel.Value = h
            totais_da_linha(ws)
            observacoes(ws, retro_itens.get(prt, []), lg)
            if len(dias_tab) > len(dias):
                lg.append(f"retroativos em {len(dias_tab) - len(dias)} coluna(s) antes do ciclo: "
                          f"{', '.join(f'{d:%d/%m}' for d in dias_tab[:len(dias_tab) - len(dias)])}")
            feitas.append(prt)

            # 3) blocos sem hora no mês ficam ocultos
            ocultos = []
            for bl in blocos:
                usa = any(any(v not in (None, "", 0) for v in ws.Range(ws.Cells(r, COL_DIA1), ws.Cells(r, soma_c - 1)).Value[0])
                          for r in bl["linhas"].values())
                ws.Rows(f"{bl['ini']}:{bl['fim']}").Hidden = not usa
                if not usa:
                    ocultos.append(bl["rot"][:30])
            log.append(f"\nPRT {prt} [aba '{ws.Name}']: {len(pedidos)} lançamentos")
            log += [f"   + {x}" for x in lg]
            log += [f"   - {f}" for f in faltas]
            if ocultos:
                log.append(f"   (ocultos, sem hora no mês: {', '.join(ocultos)})")

        # abas de PRT sem hora no mês ficam ocultas
        for ws in wb.Worksheets:
            m = re.search(r"PRT\s*_?(\d{4})", ws.Name)
            if m and ws.Visible == -1 and m.group(1) not in dados:
                ws.Visible = 0
                log.append(f"\nAba '{ws.Name}' OCULTA: PRT sem hora neste mês")
        padronizar_fontes(wb, "Consolidação", log)
        exp = []
        for prt in feitas:                     # uma planilha por PRT, só com a aba dela
            if so and prt not in so:
                continue
            print(f"planilha da PRT {prt}...", flush=True)
            exportar(xl, wb, prt, exp)
        log += (["\nPLANILHAS POR PRT:"] + [f"   {x}" for x in exp]
                + ([f"   (só a(s) PRT {', '.join(so)} regravada(s); as outras ficaram como estavam)"] if so else []))
        wb.Worksheets(1).Activate()
        wb.Save()
        wb.Close()
    os.makedirs(ARQ, exist_ok=True)            # a de várias abas sai da entrega: é o modelo do mês seguinte
    final = os.path.join(ARQ, os.path.basename(dst))
    if os.path.exists(final):
        os.remove(final)
    shutil.move(dst, final)
    log.append(f"   planilha de várias abas (modelo do mês seguinte): {final}")
    dst = final
    rel = os.path.join(C.SAIDA, "relatorio_consolidacao.txt")
    with open(rel, "w", encoding="utf-8") as f:
        f.write("\n".join(log))
    print("\n".join(log))
    print(f"\nGerado: {dst}\nRelatório: {rel}")


if __name__ == "__main__" and sys.argv[1:] == ["teste"]:     # autoteste das regras puras: python aplicar_consolidacao.py teste
    assert faixa_da_linha("=SUM(V44:AZ44)", 44, "H", "AK") == "=SUM(H44:AK44)"
    assert faixa_da_linha("=(SUM(H13:AL13))", 13, "H", "AO") == "=(SUM(H13:AO13))"
    assert faixa_da_linha("=SUM(AM15:AM18)*24", 15, "H", "AL") == "=SUM(AM15:AM18)*24"
    assert texto_retroativo([("Mecânico Montador", "Normais", "5340"), ("Mecânico Montador", "Normais", "5337")]) == \
        "Apontamentos retroativos da função de Mecânico Montador, que não estava prevista na PRT. RDOs 5337; 5340"
    assert texto_retroativo([("Topógrafo", "H. Extras 75%", "5208"), ("Topógrafo", "Ad. Noturno", "5218"),
                             ("Eletricista FC", "Ad. Noturno", "5350")]) == (
        "Apontamentos retroativos que não estavam previstos na PRT: Eletricista FC (adicional noturno); "
        "Topógrafo (horas extras e adicional noturno). RDOs 5208; 5218; 5350")
    assert tipo_rotulo("CALDEIREIRO  ADICIONAL NOTURNO") == "NOT" and tipo_rotulo("SOLDADOR HORAS 75 %") == "75"
    assert tipo_rotulo("CALDEREIRO HORAS 50") == "75" and tipo_rotulo("TECNICO DE SEGURANCA HORAS NOTURNAS") == "NOT"
    assert tipo_rotulo("CALDEREIRO") == "N" and tipo_rotulo("MECANICO MONTADOR HORAS 100") == "100"
    print("aplicar_consolidacao: ok")
elif __name__ == "__main__":
    main()
