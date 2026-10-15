# -*- coding: utf-8 -*-
"""Lê a PQ (item / unidade / quantidade contratada) do PDF da PRT aprovada.

Funciona só em PDF com texto (os digitalizados como imagem voltam vazios — aí a
quantidade contratada tem de ser digitada na aba Escopo da Base).
A linha da PRT vem como:  4.2.1.10 hh 2 Eletricista FC 8 R$ 70,00 R$ 560,00
(código, unidade, nº de pessoas/unidades, descrição, quantidade, preço, total);
a quantidade contratada é a que fecha com o total: qtd × pessoas (ou só qtd).
"""
import glob
import os
import re

from comum import norm

LINHA = re.compile(r"^(\d+(?:\.\d+)+)\s+(?:\d{1,2}\s+)?(\S+)\s+(\d+(?:[.,]\d+)?)\s+(.+?)\s+([\d.]*\d(?:,\d+)?)\s+R\$\s*([\d.]+,\d+)\s+R\$\s*([\d.]+,\d+)\s*$")


def _n(s):
    return float(s.replace(".", "").replace(",", "."))


def prt_nova(prt):
    """xlsx da PRT sem nenhum histórico (config_mes.PRT_NOVA) — ex.: 1ª medição de PRT nova."""
    import config_mes as C
    return getattr(C, "PRT_NOVA", {}).get(prt)


def pdf_da_prt(pasta, prt):
    c = [] if not pasta else [
        f for f in glob.glob(os.path.join(pasta, "*.pdf"))
        if norm(os.path.basename(f)).startswith("PRT") and prt in os.path.basename(f)
        and "BOOK" not in norm(os.path.basename(f))]
    return sorted(c)[-1] if c else prt_nova(prt)


def ler_escopo_xlsx(caminho):
    """{código: (und, qtd_contratada, preço)} da aba Formuário_PRT do xlsx da PRT — valores exatos.
    Código repetido (o mesmo item em mais de um bloco): mesmo preço -> soma as quantidades;
    preços diferentes (ex.: 6.1 fornecimentos, 'vb') -> 1 x soma dos valores, como nos BM Acumulados."""
    import openpyxl
    from comum import ADICIONAIS, eh_codigo
    ws = openpyxl.load_workbook(caminho, data_only=True)["Formuário_PRT"]
    blocos, mae = {}, None
    for r in range(12, ws.max_row + 1):
        cod, und, q, preco = (ws.cell(r, c).value for c in (3, 9, 10, 11))
        if not eh_codigo(cod) or not und or not isinstance(q, (int, float)):
            continue
        cod = str(cod).strip()
        if cod in ADICIONAIS:
            if not mae:
                continue
            cod = f"{mae}|{cod}"                 # hora extra: chave 'função|1.5.x', como no pipeline
        elif str(und).strip().lower() == "hh":
            mae = cod
        blocos.setdefault(cod, (str(und).strip(), []))[1].append((q, preco))
    out = {}
    for cod, (und, lst) in blocos.items():
        precos = {p for _, p in lst}
        if len(precos) == 1:
            out[cod] = (und, sum(q for q, _ in lst), lst[0][1])
        else:
            out[cod] = (und, 1, sum(q * p for q, p in lst))
    return out


def ler_escopo_pdf(caminho):
    """{código: (und, qtd_contratada, preço)} — vazio se o PDF não tem texto. xlsx da PRT: ler_escopo_xlsx."""
    if str(caminho).lower().endswith((".xlsx", ".xlsm")):
        return ler_escopo_xlsx(caminho)
    import pdfplumber
    out = {}
    with pdfplumber.open(caminho) as pdf:
        for pg in pdf.pages:
            for linha in (pg.extract_text() or "").splitlines():
                m = LINHA.match(linha.strip())
                if not m:
                    continue
                cod, und, n, _, q, preco, total = m.groups()
                if not re.fullmatch(r"[a-zçãéêôúí/²³ ]{1,8}", und):   # unidade de verdade (hh, m/mês, pç...)
                    continue
                n, q, preco, total = _n(n), _n(q), _n(preco), _n(total)
                for cand in (q * n, q):
                    if total and abs(cand * preco - total) <= max(1.0, total * 0.01):
                        break
                else:
                    cand = q * n
                if cod in out:                    # item repetido em outro bloco da proposta: soma
                    u0, q0, p0 = out[cod]
                    out[cod] = (u0, q0 + cand, p0) if abs(p0 - preco) < 0.005 else (u0, 1, q0 * p0 + cand * preco)
                else:
                    out[cod] = (und, cand, preco)
    return out


if __name__ == "__main__":
    import sys
    import config_mes as C
    from comum import pasta_prt
    for prt in sys.argv[1:]:
        f = pdf_da_prt(pasta_prt(C.FONTE_MEDICAO, prt, C.PASTA_DA_PRT), prt)
        print(prt, os.path.basename(f) if f else None, ler_escopo_pdf(f) if f else "")
