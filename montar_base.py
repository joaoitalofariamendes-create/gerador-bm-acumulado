# -*- coding: utf-8 -*-
"""PASSO 1 — monta a 'Base do Mês': uma linha por item a medir, por PRT.

Rode:  python montar_base.py
Abra a Base gerada em SAIDA, confira SÓ as linhas amarelas (coluna H) e salve.
Depois rode aplicar_base.py.

(teste) python montar_base.py --valores-de <pasta do mês pronta>
        preenche a coluna H com as quantidades de uma Medição já fechada.
"""
import os
import re
import sys
import zipfile

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import config_mes as C
from comum import (fonte, ADICIONAIS, TIPO_ADICIONAL, arquivo, chave_cod, eh_codigo, he_vira_normal, ler_acumulado,
                   ler_caderno, ler_medicao, norm, num, pasta_prt, pq_da_medicao, sugerir)
from escopo_pdf import ler_escopo_pdf, pdf_da_prt

# Regras de cobrança que o usuário confirmou (função apontada -> item da PRT), aplicadas antes de montar as linhas.
# 5042: todo soldador é Soldador MIG (4.2.2.22), inclusive a hora extra; o Soldador Chaparia (4.2.2.20) fica sem nada.
# 5019 (06/10, proposta válida rE só tem Soldador MIG): idem.
REMAPEAR = {"5042": {"4.2.2.20": "4.2.2.22"}, "5019": {"4.2.2.20": "4.2.2.22"}}

AMARELO = PatternFill("solid", fgColor="FFF2CC")
CINZA = PatternFill("solid", fgColor="D9D9D9")
AZUL = PatternFill("solid", fgColor="DDEBF7")
CAB = ["PRT", "Código", "Sob a função", "Descrição", "Unid.", "Mês anterior", "Caderno",
       "VALOR DO MÊS", "Situação"]
CAB_ESCOPO = ["PRT", "Código", "Sob a função", "Descrição", "Unid.", "Preço unit.",
              "QTD CONTRATADA", "Origem / por que está aqui"]


def abas_codigo(caminho):
    """Códigos das abas de item de um BM, lendo só o índice do xlsx (rápido)."""
    with zipfile.ZipFile(caminho) as z:
        x = z.read("xl/workbook.xml").decode("utf-8")
    nomes = [n.strip() for n in re.findall(r'<sheet [^>]*name="([^"]*)"', x)]
    return {n for n in nomes if eh_codigo(n)}


def escopo_da_prt(prt, lm, chaves_mes, valores_de):
    """Linhas da aba Escopo: itens do mês que não existem no Acumulado/BM do mês passado,
    e, na 1ª medição, todo o escopo do PDF da PRT."""
    acu, bm = fonte(prt, "acumulado"), fonte(prt, "bm")
    primeira = acu is None
    no_acu = {l["chave"] for l in ler_acumulado(acu)[1]} if acu else set()
    no_bm = abas_codigo(bm) if bm else set()
    pdf = pdf_da_prt(pasta_prt(C.FONTE_MEDICAO, prt, C.PASTA_DA_PRT), prt)
    esc_pdf = ler_escopo_pdf(pdf) if pdf else {}
    gab = {}
    if valores_de:
        g = arquivo(pasta_prt(valores_de, prt, C.PASTA_DA_PRT), "acumulado")
        if g:
            gab = {l["chave"]: num(l.get("contratado")) for l in ler_acumulado(g)[1]}
    pq = pq_da_medicao(fonte(prt, "medicao"))

    quero = []
    if primeira:
        quero += [(k, "1ª medição — escopo do PDF da PRT") for k in esc_pdf]
    for k in chaves_mes:
        cod, _ = chave_cod(k)
        if primeira:
            motivo = "1ª medição — item medido no mês"
        elif k not in no_acu:
            motivo = "item novo: não existe no BM Acumulado do mês passado"
        elif "|" not in k and cod not in no_bm:
            motivo = "item novo: não tem aba no BM do mês passado"
        else:
            continue
        if k not in dict(quero):
            quero.append((k, motivo))
    linhas = []
    for k, motivo in quero:
        cod, mae = chave_cod(k)
        desc, und, preco, _ = pq.get(cod, ("", "", None, None))
        if num(preco) is None and mae is None and k in esc_pdf:
            preco = esc_pdf[k][2]            # PQ sem preço (6.1 '-'): vale o valor da PRT (HE fica o % da PQ)
        if gab.get(k) is not None:
            q, origem = gab[k], "teste: quantidade do gabarito"
        elif k in esc_pdf:
            q, origem = esc_pdf[k][1], f"PRT {os.path.basename(pdf)}"
        else:
            q, origem = None, "PREENCHER: não achei a quantidade contratada"
        linhas.append([prt, cod, mae, desc, und, preco, q, f"{motivo}; {origem}"])
    return linhas


def descricoes(linhas):
    por_cod = {l["cod"]: str(l["desc"]) for l in linhas}
    out = {}
    for l in linhas:
        d = str(l["desc"])
        if l["cod"] in ADICIONAIS:
            d = f"{d[:28]}… — {por_cod.get(l['mae'], '')}"
        elif d.strip().lower().startswith(("hora produtiva", "mobiliza")) and l["cod"].count(".") >= 3:
            pai = por_cod.get(l["cod"].rsplit(".", 1)[0], "")
            d = f"{pai} — {d}"
        out[l["chave"]] = d
    return out


def main():
    valores_de = sys.argv[sys.argv.index("--valores-de") + 1] if "--valores-de" in sys.argv else None
    caderno = ler_caderno(C.CADERNO)
    os.makedirs(C.SAIDA, exist_ok=True)

    prts = set(caderno)
    for p in os.listdir(C.FONTE_MEDICAO):
        if p.upper().startswith("PRT"):
            prts.update(x for x in p.replace("-", " ").split() if x.isdigit() and len(x) == 4)

    linhas_base, pend, linhas_escopo = [], [], []
    for prt in sorted(prts):
        med = fonte(prt, "medicao")
        if med is None:
            if prt in caderno:
                pend.append((prt, "Tem apontamento no Caderno mas não achei Medição-modelo para esta PRT"))
            continue
        lm = ler_medicao(med)
        desc = descricoes(lm)
        sug, avisos = sugerir(caderno.get(prt, []), lm, prt)
        for de, para in REMAPEAR.get(prt, {}).items():
            for k in [k for k in sug if k == de or k.startswith(de + "|")]:
                novo = para + k[len(de):]
                q, _ = sug.pop(k)
                q0, _ = sug.get(novo, (None, ""))
                sug[novo] = ((q0 or 0) + (q or 0), f"regra da PRT {prt}: {de} é cobrado como {para}")
                if novo not in {l["chave"] for l in lm}:
                    pend.append((prt, f"{k} = {q} vai para {novo} (regra da PRT), mas a Medição não tem essa linha — lançar"))
        # item do Caderno sem linha na Medição: casa pela descrição da PQ -> vira item novo
        # só item ORÇADO na PRT pode ser medido: o Acumulado tem todos os itens orçados
        acu_prt = fonte(prt, "acumulado")
        orcados = ({l["cod"] for l in ler_acumulado(acu_prt)[1] if l["und"] not in (None, "", 0)}
                   if acu_prt else None)
        pq_desc = {norm(v[0]): c for c, v in pq_da_medicao(med).items()
                   if v[0] and (orcados is None or c in orcados)}
        pq_todos = {norm(v[0]): c for c, v in pq_da_medicao(med).items() if v[0]}
        novos, novos_he, he_normal = [], [], {}
        for a in avisos:
            m = re.match(r"(.+) / (.+?) = ([\d.,]+): (.*)$", a, re.S)
            if not m:
                pend.append((prt, a))
                continue
            item, tipo, q, msg = m.group(1), m.group(2), float(m.group(3).replace(",", ".")), m.group(4)
            cod = None
            if msg.startswith("nenhuma linha da Medição"):
                import difflib
                alvo = norm(item)
                for dic in (pq_desc, pq_todos):         # orçado primeiro; senão a PQ inteira (criado e sinalizado)
                    nota, d = max(((difflib.SequenceMatcher(None, alvo, d).ratio(), d) for d in dic), default=(0, None))
                    if nota >= 0.9:
                        cod = dic[d]
                        break
                if cod and he_vira_normal(prt, item, tipo):     # HE sem HE na proposta: hora normal (5007 Inspetor 100%)
                    novos.append((cod, q))
                    he_normal[cod] = he_normal.get(cod, 0) + q
                elif cod and tipo in TIPO_ADICIONAL:
                    novos_he.append((cod, TIPO_ADICIONAL[tipo], q))
                elif cod:
                    novos.append((cod, q))
            else:
                mh = re.match(r"(\S+) não tem linha (1\.5\.\d) abaixo dela", msg)
                if mh:
                    novos_he.append((mh.group(1), mh.group(2), q))
                    cod = mh.group(1)
            if not cod:
                pend.append((prt, a))

        ant = {}
        acu = fonte(prt, "acumulado")
        if acu:
            ant = {l["chave"]: l["qtd"] for l in ler_acumulado(acu)[1] if num(l["qtd"])}
        else:
            pend.append((prt, "Sem BM Acumulado do mês anterior (coluna 'Mês anterior' vazia)"))

        gab = {}
        if valores_de:
            g = arquivo(pasta_prt(valores_de, prt, C.PASTA_DA_PRT), "medicao")
            gab = {l["chave"]: l["qtd"] for l in ler_medicao(g) if num(l["qtd"])} if g else {}

        for l in lm:
            k = l["chave"]
            if not l["und"] or l["und"] == 0:
                continue
            if k not in sug and k not in gab:
                # horas só vêm do Caderno; sem apontamento = não mede
                if k not in ant or str(l["und"]).strip().lower() in ("hh", "%"):
                    continue
            q, nota = sug.get(k, (None, ""))
            if k in sug and not nota:
                valor, sit = q, "Caderno"
            elif nota.startswith("regra"):          # regra da PRT (REMAPEAR) ou regra geral (HE do operador de Munck)
                valor, sit = q, "OK — " + nota
            elif k in sug and "fração de mês" in nota:
                # veículo por mês: 1 se serviu a PRT o mês todo; fração só quando rateado entre PRTs (confirme)
                valor = 1
                sit = f"CONFERIR: {nota} (pré-preenchido com 1 = mês todo; fração só se foi rateado)"
            elif k in sug:
                valor, sit = q, "CONFERIR: " + nota
            else:
                valor, sit = None, "CONFERIR: medido no mês passado, sem apontamento no Caderno"
            if valores_de:
                valor, sit = gab.get(k), "teste: valor copiado do gabarito"
            linhas_base.append([prt, l["cod"], l["mae"], desc[k], l["und"], ant.get(k), q, valor, sit])
        if orcados is not None:
            for b in linhas_base:
                if b[0] == prt and b[7] and b[1] not in ADICIONAIS and b[1] not in orcados:
                    b[8] = f"CONFERIR: item NÃO ORÇADO na PRT (não está no BM Acumulado) — {b[8]}"
        pq_prt = pq_da_medicao(med)
        soma_novos = {}
        for cod, q in novos:                              # mesma função em mais de uma linha do Caderno: soma
            soma_novos[cod] = soma_novos.get(cod, 0) + q
        for cod, q in soma_novos.items():
            linhas_base.append([prt, cod, None, pq_prt[cod][0], pq_prt[cod][1], ant.get(cod), q, q,
                                f"CONFERIR: item sem linha na Medição — casado pela descrição da PQ ({cod}); "
                                "linha será CRIADA" + ("" if orcados is None or cod in orcados else " — NÃO ORÇADO na PRT")
                                + (f"; regra da PRT: inclui {he_normal[cod]:g} h extra(s) como hora normal" if cod in he_normal else "")])
        he_soma = {}
        for mae, ad, q in novos_he:
            he_soma[(mae, ad)] = he_soma.get((mae, ad), 0) + q
        for (mae, ad), q in he_soma.items():
            linhas_base.append([prt, ad, mae, f"{pq_prt.get(ad, ('',))[0]} — {pq_prt.get(mae, ('',))[0]}",
                                pq_prt.get(ad, (None, None))[1], ant.get(f"{mae}|{ad}"), q, q,
                                "CONFERIR: adicional sem linha na Medição — linha será CRIADA"])
        chaves_mes = [f"{b[2]}|{b[1]}" if b[2] else b[1] for b in linhas_base
                      if b[0] == prt and b[7]]
        linhas_escopo += escopo_da_prt(prt, lm, chaves_mes, valores_de)

    # ------------------------------------------------------------ grava a Base (arquivo novo)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Base"
    ws.append([f"BASE DO MÊS — BM {C.BM_NUMERO} — {C.PERIODO_INI:%d/%m/%Y} a {C.PERIODO_FIM:%d/%m/%Y}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Confira as linhas amarelas e preencha/corrija a coluna H (VALOR DO MÊS). "
               "Vazio = não mede. Para medir um item que não está aqui, acrescente uma linha no fim "
               "com PRT, Código, (Sob a função, se for hora extra) e VALOR."])
    ws.append(CAB)
    for c in ws[3]:
        c.font, c.fill = Font(bold=True), CINZA
    for lb in linhas_base:
        ws.append(lb)
        r = ws.max_row
        ws.cell(r, 8).fill = AZUL
        if str(lb[8]).startswith("CONFERIR"):
            for c in ws[r]:
                c.fill = AMARELO
    for col, w in zip("ABCDEFGHI", (7, 10, 11, 60, 13, 12, 10, 14, 70)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:I{ws.max_row}"
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    dv.add(f"H4:H{ws.max_row + 200}")
    ws.add_data_validation(dv)

    we = wb.create_sheet("Escopo")
    we.append(["ESCOPO — itens que serão CRIADOS no BM Acumulado / BM (item novo ou PRT na 1ª medição)"])
    we["A1"].font = Font(bold=True, size=13)
    we.append(["Confira a coluna G (quantidade contratada da PRT). Amarelo = não achei, preencha. "
               "Vazio fica vazio no Acumulado (coluna E) e no BM (Previsto)."])
    we.append(CAB_ESCOPO)
    for c in we[3]:
        c.font, c.fill = Font(bold=True), CINZA
    for le in linhas_escopo:
        we.append(le)
        r = we.max_row
        we.cell(r, 7).fill = AZUL
        if le[6] is None:
            for c in we[r]:
                c.fill = AMARELO
    for col, w in zip("ABCDEFGH", (7, 10, 11, 60, 9, 12, 16, 80)):
        we.column_dimensions[col].width = w
    we.freeze_panes = "A4"

    wp = wb.create_sheet("Pendências")
    wp.append(["PRT", "O que o Caderno trouxe e não tem lugar automático"])
    for c in wp[1]:
        c.font, c.fill = Font(bold=True), CINZA
    for p in pend:
        wp.append(list(p))
    wp.column_dimensions["B"].width = 120
    for row in ws.iter_rows(min_row=4):
        row[3].alignment = Alignment(wrap_text=False)

    nome = os.path.join(C.SAIDA, f"Base_do_Mes_BM{C.BM_NUMERO}{'_TESTE' if valores_de else ''}.xlsx")
    wb.save(nome)
    conferir = sum(1 for l in linhas_base if str(l[8]).startswith("CONFERIR"))
    print(f"Base gerada: {nome}")
    print(f"  {len(linhas_base)} linhas  |  {conferir} para conferir (amarelas)  |  {len(pend)} pendências")
    sem_qtd = sum(1 for l in linhas_escopo if l[6] is None)
    print(f"  Escopo: {len(linhas_escopo)} itens a criar  |  {sem_qtd} sem quantidade contratada (amarelas)")


if __name__ == "__main__":
    main()
