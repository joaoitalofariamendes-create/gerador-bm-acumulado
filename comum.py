# -*- coding: utf-8 -*-
"""Leitura das planilhas, de-para único e escrita segura via Excel.

Regras de segurança (valem para todos os scripts):
  - leitura com openpyxl (somente leitura, nunca salva por ele: o openpyxl
    descarta imagens WMF e formatação condicional ao salvar);
  - escrita SEMPRE pelo próprio Excel (COM), nunca desmescla;
  - valor vai para a célula que já existe; item/função que falta é CRIADO copiando
    uma linha/aba irmã já existente (mesma formatação e fórmulas) — nunca layout novo;
  - tudo o que foi criado, e o que não achou lugar, vai para o relatório.
"""
import glob
import os
import re
import unicodedata

import openpyxl

ADICIONAIS = {"1.5.1", "1.5.2", "1.5.3", "1.5.4"}
TIPO_ADICIONAL = {"H. Extras 75%": "1.5.2", "H. Extras 100%": "1.5.1",
                  # adicional noturno do Caderno (22h–05h): mede como 1.5.3 = função × % — antes ficava "só informativo"
                  # e a equipe da noite da 5007 (RDO 5626, 05/10/2026) saía sem o noturno
                  "Ad. Noturno": "1.5.3",
                  # pendências de ciclos anteriores (caderno, regra pontual) — inclui adicional noturno
                  "H. Extras 75% (pendência)": "1.5.2", "H. Extras 100% (pendência)": "1.5.1",
                  "Ad. Noturno (pendência)": "1.5.3"}

# Caderno -> descrição da PQ (quando o nome não é idêntico)
ALIAS = {
    "TECNICO DE PLANEJAMENTO": "TECNICO EM PLANEJAMENTO/MEDICAO",
    "ANDAIME": "MONTAGEM DE ANDAIME TUBULAR",
    "VAN 15 LUGARES": "VAN 15 LUGARES INCLUIDO O MOTORISTA",
}
# funções sem item cobrável: o retorno vem pelo equipamento que operam
SEM_ITEM = {"MONTADOR DE ANDAIME", "OPERADOR DE PTA", "OPERADOR DE CAMINHAO MUCK"}
# operador cuja hora volta pelo equipamento (item com operador incluso): a hora EXTRA dele (sábado, após o expediente)
# não tem item de HE no equipamento — entra como hora NORMAL do equipamento (usuário, 06/10/2026: Munck da 5022,
# 2 sábados = 16 h que o RDO trouxe só na mão de obra)
HE_DO_OPERADOR = {"OPERADOR DE CAMINHAO MUCK": "CAMINHAO GUINDAUTO"}
HE_TIPOS = ("H. Extras 75%", "H. Extras 100%")
# função SEM hora extra na proposta da PRT: a hora extra apontada conta como hora NORMAL (usuário, 06/10/2026, 5042 NR12:
# "Na proposta não existe hora extra para essas funções, então elas devem ser apontadas como horas normais"; e na mesma
# conversa: 5029 inteira). {PRT: {função do Caderno ou '*': tipos de HE}}
HE_COMO_NORMAL = {"5042": {"SUPERVISOR DE MECANICA": HE_TIPOS, "TECNICO DE PLANEJAMENTO": HE_TIPOS},
                  "5029": {"*": HE_TIPOS}}


def he_vira_normal(prt, func, tipo):
    regra = HE_COMO_NORMAL.get(str(prt), {})
    return tipo.replace(" (pendência)", "") in regra.get(norm(func), regra.get("*", ()))
# unidades em que o Caderno traz DIAS e a Medição cobra fração de mês
UND_FRACAO_MES = {"UND/MES"}


def norm(s):
    s = "".join(c for c in unicodedata.normalize("NFD", str(s)) if unicodedata.category(c) != "Mn")
    return " ".join(re.sub(r"[^A-Z0-9/ ]", " ", s.upper()).split())


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def eh_codigo(v):
    return v is not None and re.fullmatch(r"\d+(\.\d+)*", str(v).strip()) is not None


# ----------------------------------------------------------------- arquivos
def pasta_prt(raiz_mes, prt, overrides):
    if prt in overrides and os.path.isdir(os.path.join(raiz_mes, overrides[prt])):
        return os.path.join(raiz_mes, overrides[prt])
    achadas = [p for p in glob.glob(os.path.join(raiz_mes, "PRT*")) if os.path.isdir(p)
               and re.search(r"\b" + prt + r"\b", os.path.basename(p).replace("-", " "))]
    if len(achadas) == 1:
        return achadas[0]
    if len(achadas) > 1:
        raise ValueError(f"PRT {prt}: mais de uma pasta {achadas} — defina em PASTA_DA_PRT")
    return None


def meses_anteriores():
    """Pastas 'MES mm-aaaa - BM nnnn' da RAIZ, da mais recente para a mais antiga, antes do BM atual."""
    import config_mes as C
    out = []
    for p in glob.glob(os.path.join(C.RAIZ, "MES *")):
        m = re.search(r"BM\s*(\d{4})", os.path.basename(p))
        if m and int(m.group(1)) < int(C.BM_NUMERO):
            out.append((int(m.group(1)), p))
    return [p for _, p in sorted(out, reverse=True)]


def drive_ok(raiz, tentativas=12, espera=10):
    """O Google Drive some por instantes (04/10: o aplicar_bm não achou o BM da 5001, foi tratá-la como PRT nova e
    caiu). Antes de concluir que um arquivo não existe, a pasta do mês passado tem de listar as PRTs."""
    import time
    for _ in range(tentativas):
        if glob.glob(os.path.join(raiz, "PRT*")):
            return
        time.sleep(espera)
    raise SystemExit(f"Drive indisponível: {raiz} não lista nenhuma pasta de PRT — rode de novo quando o Drive voltar")


def fonte(prt, tipo):
    """Arquivo-fonte da PRT: pasta do mês anterior; se a PRT não foi medida lá, o último mês em
    que ela aparece. Medição que nunca existiu é montada a partir de modelo (medicao_nova)."""
    import config_mes as C
    raiz = C.FONTE_MEDICAO if tipo == "medicao" else C.FONTE_ACUMULADO
    for _ in range(2):              # 2ª volta só depois de confirmar que o Drive está listando as pastas
        f = arquivo(pasta_prt(raiz, prt, C.PASTA_DA_PRT), tipo)
        if f:
            return f
        for mes in meses_anteriores():
            try:
                f = arquivo(pasta_prt(mes, prt, C.PASTA_DA_PRT), tipo)
            except ValueError:
                continue
            if f:
                return f
        drive_ok(raiz)
    if tipo == "medicao" and (fonte(prt, "acumulado") or getattr(C, "PRT_NOVA", {}).get(prt)):
        from medicao_nova import medicao_modelo
        return medicao_modelo(prt)
    return None


def arquivo(pasta, tipo):
    """tipo: 'medicao' | 'acumulado' | 'bm'"""
    if pasta is None:
        return None
    todos = [f for f in glob.glob(os.path.join(pasta, "*.xls*")) if not os.path.basename(f).startswith("~$")]
    b = lambda f: norm(os.path.basename(f))
    if tipo == "medicao":
        c = [f for f in todos if b(f).startswith("MEDICAO")]
    elif tipo == "acumulado":
        c = [f for f in todos if "ACUMULADO" in b(f)]
    else:
        c = [f for f in todos if b(f).startswith("BM") and "ACUMULADO" not in b(f) and "CONSOLIDA" not in b(f)]
    if len(c) > 1:
        raise ValueError(f"{pasta}: vários arquivos de {tipo}: {c}")
    return c[0] if c else None


# ----------------------------------------------------------------- Caderno
def ler_caderno(caminho):
    """{prt: [(item, tipo, qtd)]} a partir da aba Medição do Caderno."""
    ws = openpyxl.load_workbook(caminho, data_only=True)["Medição"]
    blocos, atual = {}, None
    for r in range(5, ws.max_row + 1):
        a, b, c, d = (ws.cell(r, k).value for k in range(1, 5))
        if a:
            m = re.search(r"PRT_?\s*(\d{4})", str(a))
            atual = m.group(1) if m else None
            if atual:
                blocos.setdefault(atual, [])
        elif atual and b and num(d) is not None:
            blocos[atual].append((str(b).strip(), str(c).strip(), d))
    return blocos


# ----------------------------------------------------------------- linhas de item
def linhas_itens(ws, col_cod, col_desc, col_und, col_qtd, ini=12):
    """Lista de dicts {r, cod, mae, chave, desc, und, qtd, oculta}.
    chave = código, ou 'código_mãe|1.5.x' para os adicionais de hora extra."""
    out, mae, vistos = [], None, {}
    for r in range(ini, ws.max_row + 1):
        cod = ws.cell(r, col_cod).value
        if not eh_codigo(cod):
            continue
        cod = str(cod).strip()
        und = ws.cell(r, col_und).value
        if cod in ADICIONAIS:
            chave = f"{mae}|{cod}"
        else:
            chave = cod
            if str(und).strip().lower() == "hh":
                mae = cod
        oculta = bool(ws.row_dimensions[r].hidden)
        if chave in vistos:
            # código repetido: vale a linha visível; a outra fica com sufixo #L<linha>
            prev = vistos[chave]
            if prev["oculta"] and not oculta:
                prev["chave"] = f"{chave}#L{prev['r']}"
            else:
                chave = f"{chave}#L{r}"
        l = dict(r=r, cod=cod, mae=mae if cod in ADICIONAIS else None, chave=chave,
                 desc=ws.cell(r, col_desc).value or "", und=und,
                 qtd=ws.cell(r, col_qtd).value, oculta=oculta)
        vistos[l["chave"]] = l
        out.append(l)
    return out


def ler_medicao(caminho):
    ws = openpyxl.load_workbook(caminho, data_only=True)["Formuário_PRT"]
    return linhas_itens(ws, 3, 4, 9, 10)


def pq_da_medicao(caminho):
    """{código: (descrição, unidade, preço SITE, código CMS)} da aba PQ da Medição."""
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    out = {}
    for row in wb["PQ"].iter_rows(min_row=11, max_col=9, values_only=True):
        if eh_codigo(row[0]) and str(row[0]).strip() not in out:
            out[str(row[0]).strip()] = (row[1], row[2], row[5], row[8])
    wb.close()      # read_only prende o arquivo: o Save do Excel nele abria 'Salvar como' escondido e travava
    return out


def prt_info(caminho_med):
    """('PRT-GSV-DOP-B-26-5043', 'PROTEÇÃO DE PARTES ...') do cabeçalho da Medição."""
    ws = openpyxl.load_workbook(caminho_med, data_only=True)["Formuário_PRT"]
    ident = re.sub(r"\s+", "", str(ws["I6"].value or ""))
    nome = re.sub(r"^ORA-\d+-", "", " ".join(str(ws["E6"].value or "").split()))
    if not ident:
        m = re.search(r"(\d{4})", os.path.basename(os.path.dirname(caminho_med)))
        ident = f"PRT_{m.group(1)}" if m else ""
    return ident, nome


def ler_escopo(caminho_base):
    """{prt: {chave: qtd contratada}} da aba Escopo da Base (vazia se não existir)."""
    wb = openpyxl.load_workbook(caminho_base, data_only=True)
    if "Escopo" not in wb.sheetnames:
        return {}
    out = {}
    for row in wb["Escopo"].iter_rows(min_row=4, max_col=7, values_only=True):
        prt, cod, mae, _, _, _, qtd = row
        if prt and cod:
            chave = f"{str(mae).strip()}|{str(cod).strip()}" if mae else str(cod).strip()
            out.setdefault(str(prt).strip(), {})[chave] = num(qtd)
    return out


def chave_cod(chave):
    """'4.2.2.20|1.5.2' -> ('1.5.2', '4.2.2.20');  '5.2.7' -> ('5.2.7', None)"""
    if "|" in chave:
        mae, cod = chave.split("|")
        return cod, mae
    return chave, None


def com_funcao(faltam, existentes):
    """Hora extra de função sem linha (5029, BM 5031: Soldador TIG só no sábado 03/10) — a função entra com 0 h,
    antes dela, para a HE ter linha e preço: item apontado sem linha é criado, mesmo fora da PRT."""
    tem = set(existentes) | {k for k, _ in faltam}
    for k, _ in list(faltam):
        mae = chave_cod(k)[1]
        if mae and mae not in tem:
            faltam.append((mae, 0))
            tem.add(mae)
    return faltam


def ordem(cod):
    return tuple(int(x) for x in str(cod).split("."))


def aba_acumulado(wb):
    """A aba do BM Acumulado que tem as colunas 'Acum. Anterior / No Mês / Acum. Atual'."""
    cands = [ws for ws in wb.worksheets if norm(ws.cell(8, 11).value or "") == "NO MES"]
    vis = [ws for ws in cands if ws.sheet_state == "visible"]
    return (vis or cands or [None])[0]


def ler_acumulado(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = aba_acumulado(wb)
    linhas = linhas_itens(ws, 2, 3, 4, 11)
    for l in linhas:
        l["contratado"] = ws.cell(l["r"], 5).value
        l["anterior"] = ws.cell(l["r"], 10).value
        l["atual"] = ws.cell(l["r"], 12).value
    return ws.title, linhas


# ----------------------------------------------------------------- sugestão a partir do Caderno
def sugerir(itens_caderno, linhas_med, prt=None):
    """Devolve ({chave: (qtd, nota)}, [avisos]). Casa pelo nome da PQ, só em linhas
    que existem na Medição da PRT. prt: para as regras da PRT (HE_COMO_NORMAL)."""
    por_desc = {}
    for i, l in enumerate(linhas_med):
        if l["cod"] in ADICIONAIS:
            continue
        alvo = l
        if not l["und"] or l["und"] == 0:   # título de equipamento -> filho "Hora produtiva"
            filhos = [x for x in linhas_med[i + 1:i + 4]
                      if x["cod"].startswith(l["cod"] + ".") and x["und"] not in (None, 0, "vb")]
            if not filhos:
                continue
            alvo = filhos[0]
        por_desc.setdefault(norm(l["desc"]), alvo)

    por_cod = {l["cod"]: l for l in linhas_med if l["cod"] not in ADICIONAIS}

    def achar(nome):
        m = re.match(r"\s*(\d+(?:\.\d+)+)\s*-", str(nome))   # item do Caderno já com o código: '2.13.7.14-TUBOS...'
        if m:
            return por_cod.get(m.group(1))
        n = norm(nome)
        n = ALIAS.get(n, n)
        if n in por_desc:
            return por_desc[n]
        c = [v for k, v in por_desc.items() if k.startswith(n[:35]) or n.startswith(k[:35])]
        return c[0] if len(c) == 1 else None

    sug, avisos = {}, []
    for item, tipo, q in itens_caderno:
        if norm(item) in SEM_ITEM:
            equip = HE_DO_OPERADOR.get(norm(item))
            if equip and tipo.replace(" (pendência)", "") in HE_TIPOS and q:
                cam = next((i for i, _, _ in itens_caderno if norm(i).startswith(equip)), None)
                l = achar(cam) if cam else next((v for k, v in por_desc.items() if k.startswith(equip)), None)
                if l is None:
                    avisos.append(f"{item} / {tipo} = {q}: hora extra do operador sem a linha do equipamento na Medição")
                    continue
                qa, na = sug.get(l["chave"], (0, ""))
                nota = f"regra: + {q:g} h extra(s) do operador como hora normal (equipamento sem hora extra)"
                sug[l["chave"]] = ((qa or 0) + q, "; ".join(x for x in (na, nota) if x))
            continue
        l = achar(item)
        if l is None:
            avisos.append(f"{item} / {tipo} = {q}: nenhuma linha da Medição com esse nome")
            continue
        nota = regra = ""
        if he_vira_normal(prt, item, tipo) and q:
            tipo, regra = "Normais", f"regra da PRT: + {q:g} h extra(s) como hora normal (proposta sem hora extra para a função)"
        if tipo.replace(" (pendência)", "") in ("Normais", "Equipamento/Volume"):
            alvo = l
        elif tipo in TIPO_ADICIONAL:
            ad = [x for x in linhas_med if x["mae"] == l["cod"] and x["cod"] == TIPO_ADICIONAL[tipo]]
            if not ad:
                avisos.append(f"{item} / {tipo} = {q}: {l['cod']} não tem linha {TIPO_ADICIONAL[tipo]} abaixo dela")
                continue
            alvo = ad[0]
        else:
            avisos.append(f"{item} / {tipo} = {q}: tipo não é medido (só informativo)")
            continue
        if norm(alvo["und"] or "") in UND_FRACAO_MES:
            nota = f"Caderno traz {q} dia(s); a unidade é fração de mês — informe a fração"
            q = None
        elif alvo["oculta"]:
            nota = "linha está OCULTA na Medição — confirme se mede neste item"
        qa, na = sug.get(alvo["chave"], (0, ""))
        sug[alvo["chave"]] = (None if q is None else (qa or 0) + q, "; ".join(x for x in (na, regra, nota) if x))
    return sug, avisos


# ----------------------------------------------------------------- Excel (COM)
class Excel:
    """with Excel() as xl: wb = xl.abrir(caminho) ... wb.Save(); wb.Close()"""

    def __enter__(self):
        import threading
        import win32com.client
        import win32process
        self.app = win32com.client.DispatchEx("Excel.Application")
        self.app.Visible = False
        self.app.DisplayAlerts = False
        self.app.AskToUpdateLinks = False
        self._fim = threading.Event()
        pid = win32process.GetWindowThreadProcessId(self.app.Hwnd)[1]
        threading.Thread(target=_vigia, args=(pid, self._fim), daemon=True).start()
        return self

    def abrir(self, caminho, leitura=False):
        wb = self.app.Workbooks.Open(os.path.abspath(caminho), UpdateLinks=0)
        if wb is None:      # outro arquivo de MESMO NOME já aberto nesta instância: o Excel não abre e devolve nada
            raise RuntimeError(f"Excel não abriu {caminho} — já há um arquivo com esse nome aberto")
        if wb.ReadOnly and not leitura:
            # arquivo aberto por alguém: o Save seria descartado sem erro
            wb.Close(SaveChanges=False)
            raise PermissionError(f"Arquivo aberto em outro Excel (somente leitura) — feche e rode de novo: {caminho}")
        return wb

    def __exit__(self, *a):
        self._fim.set()
        # limpeza: se o Excel já caiu/desconectou depois do trabalho salvo, não transforma sucesso em erro
        try:
            for wb in list(com(lambda: self.app.Workbooks)):
                com(lambda: wb.Close(SaveChanges=False))
        except Exception:
            pass
        try:
            com(lambda: self.app.Quit())
        except Exception:
            pass


def _vigia(pid, fim, espera=15):
    """O Excel oculto parado numa janela ('Salvar como', aviso…) trava o script sem erro e sem ninguém ver
    (04/10: 82 min parado). Janela aberta há 'espera' s é fechada: o comando do Excel falha e o script para
    com erro — o usuário pediu decisão automática em vez de esperar alguém clicar."""
    import sys
    import time
    import win32con
    import win32gui
    import win32process
    desde = {}
    while not fim.wait(3):
        abertas = []

        def ver(h, _):
            if (win32gui.IsWindowVisible(h) and win32gui.GetClassName(h) == "#32770"
                    and win32process.GetWindowThreadProcessId(h)[1] == pid):
                abertas.append(h)
            return True
        try:
            win32gui.EnumWindows(ver, None)
        except Exception:
            continue
        desde = {h: desde.get(h, time.time()) for h in abertas}
        for h, t in desde.items():
            if time.time() - t >= espera:
                print(f"!! Excel parado na janela '{win32gui.GetWindowText(h)}' há {espera}s — fechada (o passo falha)",
                      file=sys.stderr, flush=True)
                win32gui.PostMessage(h, win32con.WM_CLOSE, 0, 0)


AJUSTE = re.compile(r"\s*[+-]\s*\d+(?:[.,]\d+)?\s*$")   # número digitado no fim de uma fórmula


def tirar_ajustes(ws, onde, log):
    """Total com número digitado depois da fórmula (ex.: '=SUM(N12:N43)+7438.65', '=L32+L57-0.01')
    volta a ser fórmula pura; o relatório mostra quanto o ajuste escondia. Conta só de números
    ('=96+132') vira o valor dela (228) — o número não muda, some a entrada manual."""
    rng = ws.UsedRange
    r0, c0 = rng.Row, rng.Column
    forms = rng.Formula
    achados, contas = [], []
    for i, linha in enumerate(forms):
        for j, f in enumerate(linha):
            if not (isinstance(f, str) and f.startswith("=") and AJUSTE.search(f)):
                continue
            cel = ws.Cells(r0 + i, c0 + j)
            if not re.search(r"[A-Za-z]", f[1:]):
                v = cel.Value2
                cel.Value = v
                contas.append(f"{cel.Address.replace('$', '')} {f} -> {v:g}")
                continue
            antes = cel.Value2
            cel.Formula = AJUSTE.sub("", f)
            achados.append((cel, f, antes))
    if achados:
        ws.Calculate()
        for cel, f, antes in achados:
            depois = cel.Value2
            log.append(f"   !! {onde} {ws.Name}!{cel.Address.replace('$', '')}: tinha valor digitado ('{f}') — "
                       f"retirado; valor {antes} -> {depois} (diferença {(antes or 0) - (depois or 0):+.10g}) — conferir a origem")
    if contas:
        log.append(f"   {onde} {ws.Name}: conta digitada no lugar de valor virou o valor: {', '.join(contas)}")


def com(fn, tentativas=20):
    """Chamada COM com nova tentativa quando o Excel está ocupado ('A chamada foi rejeitada pelo chamado')."""
    import time
    import pywintypes
    for k in range(tentativas):
        try:
            return fn()
        except pywintypes.com_error as e:
            if e.hresult != -2147418111 or k == tentativas - 1:
                raise
            time.sleep(0.5)


def padronizar_fontes(wb, onde, log):
    """Em cada aba, célula com fonte fora da predominante passa para a predominante (só o NOME da fonte;
    tamanho, negrito e cor ficam). Aba já uniforme não é percorrida; linha uniforme é lida de uma vez
    (célula a célula só nas linhas com fontes misturadas — no BM de muitas abas eram milhares de chamadas)."""
    from collections import Counter
    trocas = 0
    for ws in wb.Worksheets:
        rng = com(lambda: ws.UsedRange)
        if com(lambda: rng.Font.Name) is not None:           # uma fonte só na aba inteira
            continue
        r0, c0, vals = rng.Row, rng.Column, com(lambda: rng.Value)
        if not isinstance(vals, tuple):
            continue
        ncol = len(vals[0])
        itens = []                              # (alvo COM, nome da fonte, nº de células com valor)
        for i, row in enumerate(vals):
            cheias = [j for j, v in enumerate(row) if v not in (None, "")]
            if not cheias:
                continue
            linha = ws.Range(ws.Cells(r0 + i, c0), ws.Cells(r0 + i, c0 + ncol - 1))
            n = com(lambda: linha.Font.Name)
            if n is not None:
                itens.append((linha, n, len(cheias)))
                continue
            for j in cheias:
                c = ws.Cells(r0 + i, c0 + j)
                itens.append((c, com(lambda: c.Font.Name), 1))
        cont = Counter()
        for _, n, k in itens:
            if n:
                cont[n] += k
        if not cont:
            continue
        principal = cont.most_common(1)[0][0]
        for alvo, n, k in itens:
            if n != principal:
                com(lambda: setattr(alvo.Font, "Name", principal))
                trocas += k
    if trocas:
        log.append(f"   {onde}: {trocas} célula(s) com fonte diferente da predominante da aba — padronizadas")


def data_excel(d):
    """Data como número do Excel (dia cheio). datetime gravado pelo COM chega com a hora do fuso (03:00)."""
    import datetime as _dt
    return float((d - _dt.date(1899, 12, 30)).days)


def alinhar_cabecalho(ws, log):
    """Medição: linha 6 deslocada 1 coluna à direita dos rótulos da linha 5 (BM 5030 da 5008: D6:E6 e F6:G6
    mescladas, PRT em J6, data em K6) volta ao lugar — senão PRT/data/nome saem nas células erradas."""
    if not (ws.Range("E5").MergeArea.Address == "$E$5:$F$5" and ws.Range("D6").MergeArea.Address == "$D$6:$E$6"
            and ws.Range("F6").MergeArea.Address == "$F$6:$G$6"):
        return
    ws.Range("D6").MergeArea.UnMerge()
    ws.Range("F6").MergeArea.UnMerge()
    ws.Range("F6:M6").Cut(ws.Range("E6"))
    ws.Range("E6:F6").Merge()
    log.append("   Medição: cabeçalho (linha 6) estava 1 coluna à direita dos rótulos — realinhado")


def data_br(cel):
    """Data com formato americano (m/d/aaaa: 05/10 aparecia '10/5/2026') passa a dd/mm/aaaa — só o formato."""
    if isinstance(cel.Value2, float) and re.match(r"m+/d+", cel.NumberFormatLocal):
        cel.NumberFormatLocal = "dd/mm/aaaa"
        return True
    return False


def casas_valor(x):
    """Casas decimais do valor, até 4 (2,8264 -> 4; 0,5 -> 1; 7 -> 0)."""
    return next((n for n in range(5) if abs(round(x, n) - x) < 1e-9), 4)


def formato_local_com(fmt, n):
    """Formato local do Excel (pt-BR '#.##0,00') mostrando n casas; None se não reconhecer."""
    if re.search(r"#\.##0,0+", fmt):
        return re.sub(r"#\.##0,0+", "#.##0," + "0" * n, fmt)
    if "#.##0" in fmt:
        return fmt.replace("#.##0", "#.##0," + "0" * n)
    if re.fullmatch(r"0(,0+)?", fmt):
        return "0," + "0" * n
    return None


def mostrar_casas(ws, endereco, colunas):
    """Número cujo formato esconde casas (5,1914 aparecendo 5,191) passa a mostrar as do valor (até 4) —
    só o formato muda. Percentual e 'Geral' ficam. Devolve o que mudou ('U20: 2 -> 4 casas')."""
    rng = ws.Range(endereco)
    r0, c0 = rng.Row, rng.Column
    vals = rng.Value2 if rng.Count > 1 else ((rng.Value2,),)
    feito = []
    for i, linha in enumerate(vals):
        for j, v in enumerate(linha):
            if c0 + j not in colunas or not isinstance(v, float) or abs(v) >= 1e6:
                continue
            n = casas_valor(v)
            cel = ws.Cells(r0 + i, c0 + j)
            f = cel.NumberFormatLocal
            s = re.sub(r'"[^"]*"|\\.|_.|\*.', "", f.split(";")[0])
            if not n or "%" in s or not re.search(r"[0#]", s):
                continue
            m = re.search(r",([0#]+)", s)
            d = len(m.group(1)) if m else 0
            novo = formato_local_com(f, n) if d < n else None
            if novo:
                cel.NumberFormatLocal = novo
                feito.append(f"{cel.Address.replace('$', '')}: {d} -> {n} casas")
    return feito


PQ_MEDICAO = "PQ!$A$12:$K$3000"


def procv_pq(ws, r):
    """Linha de item da Medição: descrição (D), unidade (I) e preço (K) pela PQ da própria planilha.
    Linha copiada de outra pode trazer preço fixo (ex.: 6.1) ou PQ de outro arquivo."""
    for col, n in (("D", 2), ("I", 3), ("K", 6)):
        ws.Range(f"{col}{r}").Formula = f'=IFERROR(VLOOKUP(C{r},{PQ_MEDICAO},{n},0),"")'


EXTERNO = re.compile(r"'[^']*\[[^\]]+\][^']*'!\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?\d+|\[\d+\][^!'(),]+!\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?\d+")


def pq_interna(wb, ws, log, onde="Medição"):
    """PROCV que busca a PQ de OUTRO arquivo (proposta antiga no Drive) passa a usar a PQ da própria planilha
    — só fica a troca se o valor não mudar; o vínculo que sobrar sem uso é removido."""
    usados, n = set(), 0
    rng = ws.UsedRange
    forms = rng.Formula
    for i, row in enumerate(forms):
        for j, f in enumerate(row):
            if not (isinstance(f, str) and f.startswith("=") and "[" in f and "VLOOKUP" in f.upper()):
                continue
            cel = ws.Cells(rng.Row + i, rng.Column + j)
            novo = EXTERNO.sub(PQ_MEDICAO, f)
            if novo == f:
                continue
            antes = cel.Value2
            cel.Formula = novo
            if cel.Value2 != antes:
                cel.Formula = f
                log.append(f"   {onde}: {cel.Address.replace('$', '')} busca PQ de outro arquivo e a PQ da planilha dá outro valor — mantida")
                continue
            usados.update(re.findall(r"\[([^\]]+)\]", f))
            n += 1
    if n:
        log.append(f"   {onde}: {n} fórmula(s) com PQ de outro arquivo passam a usar a PQ da planilha (mesmo valor)")
        for l in (wb.LinkSources(1) or ()):
            if any(u in l for u in usados):
                try:
                    wb.BreakLink(l, 1)
                except Exception:
                    pass


def acertar_totais_medicao(ws, log):
    """Medição (Formuário_PRT): cada 'Custo do Item' soma o bloco até a linha anterior a ele (linha criada
    no fim do bloco ficava fora do SUM) e o 'Custo Total da PRT' soma todos os 'Custo do Item'."""
    L = ws.Range("L12:L400").Formula
    C_ = ws.Range("C12:C400").Value
    K_ = ws.Range("K12:K400").Value
    subs, total = [], None
    for i, ((f,), (c,), (k,)) in enumerate(zip(L, C_, K_)):
        r, rot = 12 + i, norm(f"{k or ''} {c or ''}")
        if not str(f or "").startswith("="):
            continue
        if "CUSTO DO ITEM" in rot:
            subs.append(r)
        elif "CUSTO TOTAL" in rot:
            total = r
            break
    trocas = []
    for r in subs:
        f = str(ws.Range(f"L{r}").Formula)
        m = re.fullmatch(r"=SUM\(L\$?(\d+):L\$?(\d+)\)", f.replace(" ", ""))
        if m and int(m.group(2)) != r - 1:
            novo = f"=SUM(L{m.group(1)}:L{r - 1})"
            ws.Range(f"L{r}").Formula = novo
            trocas.append(f"L{r} {f} -> {novo}")
    if total and subs:
        f = str(ws.Range(f"L{total}").Formula)
        refs = {int(x) for x in re.findall(r"L\$?(\d+)", f)}
        if not set(subs) <= refs:
            novo = "=" + "+".join(f"L{s}" for s in subs)
            ws.Range(f"L{total}").Formula = novo
            trocas.append(f"L{total} (Custo Total) {f} -> {novo}")
    if trocas:
        log.append(f"   Medição: totais acertados para cobrir todas as linhas: {'; '.join(trocas)}")


def escrever(ws, r, c, valor, log):
    """Escreve valor numa célula existente via COM. Recusa fórmula e célula mesclada fora da âncora."""
    cel = ws.Cells(r, c)
    if cel.HasFormula and re.search(r"[A-Za-z]", cel.Formula[1:]):
        # fórmula de verdade (referência/função); conta digitada tipo "=3,4+1,51" pode ser trocada
        log.append(f"   L{r}C{c}: tem fórmula ({cel.Formula}) — não sobrescrevi")
        return False
    if cel.MergeCells and cel.MergeArea.Cells(1, 1).Address != cel.Address:
        log.append(f"   L{r}C{c}: é parte de mesclagem — não escrevi")
        return False
    cel.Value = valor
    return True


if __name__ == "__main__":      # autoteste das regras puras (sem Excel): python comum.py
    assert [casas_valor(x) for x in (2.8264, 0.5, 7.0, 1913.8332, 5.191000000000001)] == [4, 1, 0, 4, 3]
    assert formato_local_com('_-* #.##0,00_-;-* #.##0,00_-', 4) == '_-* #.##0,0000_-;-* #.##0,0000_-'
    assert formato_local_com("0", 3) == "0,000" and formato_local_com("Geral", 2) is None
    _lm = [{"cod": "5.1.12", "und": None, "desc": "Caminhão Guindauto 10 ton(Incluindo Operador) com joystick remoto",
            "chave": "5.1.12", "mae": None, "oculta": False},
           {"cod": "5.1.12.1", "und": "h", "desc": "Hora produtiva", "chave": "5.1.12.1", "mae": None, "oculta": False}]
    _s, _ = sugerir([("OPERADOR DE CAMINHAO MUCK", "Normais", 175), ("OPERADOR DE CAMINHAO MUCK", "H. Extras 75%", 16),
                     ("Caminhão Guindauto 10 ton(Incluindo Operador) com joystick remoto", "Equipamento/Volume", 175)], _lm)
    assert _s["5.1.12.1"][0] == 191 and _s["5.1.12.1"][1].startswith("regra"), _s
    _lm2 = [{"cod": "4.2.2.2", "und": "hh", "desc": "Supervisor de Mecânica", "chave": "4.2.2.2", "mae": None, "oculta": False},
            {"cod": "1.5.2", "und": "%", "desc": "Adicional Horas Extras", "chave": "4.2.2.2|1.5.2", "mae": "4.2.2.2", "oculta": False}]
    _s, _ = sugerir([("Supervisor de Mecânica", "Normais", 350), ("Supervisor de Mecânica", "H. Extras 75%", 12)], _lm2, "5042")
    assert _s == {"4.2.2.2": (362, "regra da PRT: + 12 h extra(s) como hora normal (proposta sem hora extra para a função)")}, _s
    assert he_vira_normal("5029", "Caldeireiro", "H. Extras 75% (pendência)")
    assert not he_vira_normal("5008", "Caldeireiro", "H. Extras 75%")
    assert not he_vira_normal("5042", "Mecânico Montador", "H. Extras 75%")
    assert com_funcao([("4.2.2.19|1.5.2", 11)], {"4.2.2.20"}) == [("4.2.2.19|1.5.2", 11), ("4.2.2.19", 0)]
    assert com_funcao([("4.2.2.20|1.5.2", 8)], {"4.2.2.20"}) == [("4.2.2.20|1.5.2", 8)]
    assert com_funcao([("4.2.2.19|1.5.2", 3), ("4.2.2.19", 5)], set()) == [("4.2.2.19|1.5.2", 3), ("4.2.2.19", 5)]
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        try:
            drive_ok(d, tentativas=1, espera=0)
            raise AssertionError("pasta sem PRT passou como Drive ok")
        except SystemExit:
            pass
        os.mkdir(os.path.join(d, "PRT 5001"))
        drive_ok(d, tentativas=1, espera=0)
    print("comum: ok")
