# -*- coding: utf-8 -*-
"""Confere as planilhas de Consolidação de Horas por PRT (SAIDA\\PRT xxxx\\*CONSOLIDA*.xlsx) contra o Caderno, DIA A DIA:
cada hora do Caderno (aba Consolidação + Retroativos) tem de estar na linha certa (função/normal/75/100/noturno) e no
dia certo; e nada a mais na grade. Também: data de entrega, período, retroativos citados nas observações.
Rode:  python conferir_consolidacao.py   -> SAIDA\\conferencia_consolidacao.txt"""
import datetime as dt
import glob
import os
import re
import shutil
import tempfile

import openpyxl

import config_mes as C
from aplicar_consolidacao import (APELIDOS, FORA_DA_CONSOLIDACAO, SEM_CONSOLIDACAO, TIPO, achar_grupo, com_retroativos,
                                  horas, ler_caderno_consolidacao, ler_retroativos, por_rotulo, tipo_rotulo)
from comum import SEM_ITEM, norm


def grade(ws):
    """({rot: {tipo: linha}} com a mesma escolha do gerador, {date: coluna}, linha DISCRIMINAÇÃO, todas as linhas)."""
    datas = {c.value.date(): c.column for c in ws[11] if isinstance(c.value, dt.datetime)}
    disc = next(r for r in range(13, ws.max_row + 1) if norm(ws.cell(r, 2).value or "").startswith("DISCRIMINA"))
    lista, atual = [], None
    for r in range(13, disc):
        b = ws.cell(r, 2).value
        if not b:
            continue
        n, t = norm(b), tipo_rotulo(norm(b))
        if t == "N" or atual is None:
            atual = {"rot": n, "linhas": {t: r}, "oculta": bool(ws.row_dimensions[r].hidden)}
            lista.append(atual)
        else:
            atual["linhas"].setdefault(t, r)
    blocos = {k: bl["linhas"] for k, bl in por_rotulo(lista).items()}
    return blocos, datas, disc, [r for bl in lista for r in bl["linhas"].values()]


def conferir_arquivo(f, por_func, itens):
    """Uma planilha de Consolidação (1ª aba) x horas do Caderno ({(função, tipo): {data: h}}) e RDOs de retroativo
    [(função, tipo, rdo)] -> (problemas, nº de dias, nº de retroativos, células lançadas)."""
    copia = shutil.copy2(f, os.path.join(tempfile.mkdtemp(), "c.xlsx"))    # lê cópia: aberta no Excel também serve
    ws = openpyxl.load_workbook(copia, data_only=True).worksheets[0]
    blocos, datas, disc, todas = grade(ws)
    esperado, err = {}, []
    for (func, tipo), por_dia in por_func.items():
        n = norm(func)
        if (tipo in SEM_CONSOLIDACAO or (n in APELIDOS and not APELIDOS[n]) or n in SEM_ITEM
                or n.startswith(FORA_DA_CONSOLIDACAO) or TIPO.get(tipo) is None):
            continue
        g = achar_grupo(blocos, func)
        if g is None or TIPO[tipo] not in blocos[g]:
            err.append(f"{func} / {tipo}: sem linha na planilha ({sum(por_dia.values()):g})")
            continue
        for d, h in por_dia.items():
            k = (blocos[g][TIPO[tipo]], d)
            esperado[k] = esperado.get(k, 0) + h
    achado = {}
    for r in todas:                       # todas as linhas: valor em linha errada também aparece
        for d, c in datas.items():
            v = ws.cell(r, c).value
            if v not in (None, "", 0) and not isinstance(v, str):
                achado[(r, d)] = horas(v)
    for k in sorted(set(esperado) | set(achado)):
        e, a = esperado.get(k, 0), achado.get(k, 0)
        if isinstance(ws.cell(k[0], datas[k[1]]).value, (int, float)) and not isinstance(ws.cell(k[0], datas[k[1]]).value, bool):
            a = ws.cell(k[0], datas[k[1]]).value          # quantidade (equipamento) ou hora gravada como número
            if e and abs(a - e) > 1e-6 and abs(a * 24 - e) < 1e-6:
                err.append(f"linha {k[0]} {ws.cell(k[0], 2).value} {k[1]:%d/%m}: hora gravada como número ({a})")
                continue
        if abs(e - a) > 1e-6:
            err.append(f"linha {k[0]} {str(ws.cell(k[0], 2).value)[:30]} {k[1]:%d/%m}: Caderno {e:g} x planilha {a:g}")
    # cabeçalho e observações
    cab = " ".join(str(c.value) for row in ws.iter_rows(min_row=6, max_row=9) for c in row if c.value is not None)
    if f"{C.PERIODO_INI:%d/%m/%Y} A {C.PERIODO_FIM:%d/%m/%Y}" not in cab:
        err.append("período do cabeçalho não é o do ciclo")
    if getattr(C, "DATA_BM", None) and not any(isinstance(c.value, dt.datetime) and c.value.date() == C.DATA_BM
                                               for row in ws.iter_rows(min_row=6, max_row=10) for c in row):
        err.append(f"data de entrega diferente de {C.DATA_BM:%d/%m/%Y}")
    obs = " ".join(str(ws.cell(r, 3).value or "") for r in range(disc, disc + 15))
    for _, _, rdo in itens:
        if rdo and rdo not in obs:
            err.append(f"RDO {rdo} do retroativo não está nas observações")
    return err, len(datas), sum(1 for d in datas if d < C.PERIODO_INI), len(achado)


def main():
    dados, _ = ler_caderno_consolidacao()
    retro, itens, _ = ler_retroativos()
    dados, _ = com_retroativos(dados, retro)          # mesmas regras do gerador (Munck, HE que vira normal)
    rel, total_err = [f"CONFERÊNCIA DA CONSOLIDAÇÃO POR PRT — {C.SAIDA}"], 0
    for f in sorted(glob.glob(os.path.join(C.SAIDA, "PRT *", "*CONSOLIDA*.xlsx"))):
        if os.path.basename(f).startswith("~$"):          # trava do Excel (planilha aberta pelo usuário)
            continue
        prt = re.search(r"PRT (\d{4})\.xlsx$", f).group(1)
        err, n_dias, n_ret, n_cel = conferir_arquivo(f, dados.get(prt, {}), itens.get(prt, []))
        rel.append(f"\nPRT {prt}: {'OK' if not err else f'{len(err)} problema(s)'} — {n_dias} dias "
                   f"({n_ret} retroativo(s)), {n_cel} célula(s) lançadas")
        rel += [f"   !! {x}" for x in err]
        total_err += len(err)
    rel.append(f"\nTOTAL: {total_err} problema(s)")
    open(os.path.join(C.SAIDA, "conferencia_consolidacao.txt"), "w", encoding="utf-8").write("\n".join(rel))
    print("\n".join(rel))


if __name__ == "__main__":
    main()
