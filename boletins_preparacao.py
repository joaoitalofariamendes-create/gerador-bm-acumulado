# -*- coding: utf-8 -*-
"""Boletins da Preparação de Cargas (contrato 4600000002) na pasta BM Manutenção: Equipe Fixa, Parada, Caldeiraria.

Horas do mês = aba Medição do Caderno (blocos 'CLIENTE - PREPARAÇÃO DE CARGAS - <EQUIPE>'): Fixa = RDOs do Encarregado A,
Caldeiraria = RDOs do Encarregado B ('ATIVIDADES: CALDEIRARIA'), Parada = projeto '... - PARADA'. 'H. Extras 75%' do Caderno
vai na linha HE-50% (1ª faixa de HE deste contrato), como no BM de agosto.
Rolagem POR EQUIPE (decisão do usuário no BM 5031): J = J+K e N35 = N35+P35 do boletim do mês passado da mesma equipe.
Cabeçalho: N.º BM (T2), Data BM (W2 = config_mes.DATA_BM), período (U3). FOTOS.xlsx: fotos da subpasta Fotos na grade
que já existe (fotos_bm; mais dias que vagas = dias espaçados do 1º ao último) e o N.º BM no cabeçalho.
O mês passado é lido do backup (_arquivo\\backup_boletins_preparacao_BMnnnn, feito na 1ª rodada): rodar de novo não rola 2x.
Uso: python boletins_preparacao.py <nº Fixa> <nº Parada> <nº Caldeiraria>     (BM 5031: 037 038 039)
     python boletins_preparacao.py teste
"""
import glob
import os
import re
import shutil
import sys
import tempfile
from collections import defaultdict

import openpyxl

import config_mes as C
from comum import Excel, data_excel, escrever
from fotos_bm import LOGO_ALTURA
from fotos_bm import main as preencher_fotos

AQUI = os.path.dirname(os.path.abspath(__file__))
PASTA = os.path.join(AQUI, "BM Manutenção")
BACKUP = os.path.join(AQUI, "_arquivo", f"backup_boletins_preparacao_BM{C.BM_NUMERO}")
# equipe no Caderno -> (subpasta, rótulo do N.º BM no boletim, rótulo no FOTOS)
EQUIPES = {"EQUIPE FIXA": ("Equipe Fixa", "EQ. FIXA", "Eq. Fixa"),
           "PARADA": ("Equipe Parada", "EQ. PARADA", "Eq. Parada"),
           "CALDEIRARIA": ("Equipe Caldeiraria", "EQ. CALDEIRARIA", "Eq. Caldeiraria")}
FUNCAO = {"Encarregado de Mecânica": 0, "Mecânico Montador": 1, "Ajudante de Mecânica": 1,
          "Soldador Chaparia": 2, "Caldeireiro": 3, "Técnico em Segurança do Trabalho": 4}
TIPO = {"Normais": 10, "H. Extras 75%": 15, "H. Extras 100%": 20, "Ad. Noturno": 25}
LINHAS = list(range(10, 30)) + [31, 32]          # 31/32 = mobilização/desmobilização (rolam, não medem)
ABA = "Modelo BM em branco"


def horas_caderno():
    """{equipe: {linha do boletim: horas}} da aba Medição do Caderno."""
    ws = openpyxl.load_workbook(shutil.copy2(C.CADERNO, tempfile.mktemp(suffix=".xlsx")), data_only=True)["Medição"]
    out, eq = defaultdict(lambda: defaultdict(float)), None
    for row in ws.iter_rows(values_only=True):
        v = [x for x in row if x not in (None, "")]
        if not v:
            continue
        proj = next((x for x in v if isinstance(x, str) and x.upper().startswith("CLIENTE")), None)
        if proj:                                   # cabeçalho de bloco (projeto)
            eq = proj.split(" - ")[-1].strip().upper() if "PREPARA" in proj.upper() else None
            continue
        if eq and len(v) >= 3 and v[1] in TIPO:
            if eq not in EQUIPES or v[0] not in FUNCAO:
                raise SystemExit(f"!! Caderno: equipe/função sem linha no boletim: {eq} / {v[0]} — acrescentar no script")
            out[eq][TIPO[v[1]] + FUNCAO[v[0]]] += float(v[2])
    return out


def dias_espalhados(dias, n):
    """Mais dias que vagas: n dias espaçados do 1º ao último (o rodízio do fotos_bm pegaria só os primeiros)."""
    if len(dias) <= n:
        return dias
    return [dias[round(k * (len(dias) - 1) / (n - 1))] for k in range(n)] if n > 1 else dias[:1]


def boletim(xl, f, orig, horas, num, rot, log):
    wf = openpyxl.load_workbook(orig)[ABA]
    wv = openpyxl.load_workbook(orig, data_only=True)[ABA]
    p_ant = sum((wf.cell(r, 6).value or 0) * (wf.cell(r, 11).value or 0) for r in LINHAS)
    if abs(p_ant - (wv["P35"].value or 0)) > 1e-6:
        raise SystemExit(f"!! {f}: P35 do mês passado ({wv['P35'].value}) ≠ soma F×K ({p_ant}) — conferir antes")
    n_ant = wf["N35"].value
    if not isinstance(n_ant, (int, float)):
        raise SystemExit(f"!! {f}: N35 do mês passado não é número ({n_ant!r})")
    wb = xl.abrir(f)
    ws = wb.Worksheets(ABA)
    for r in LINHAS:
        j, k = wf.cell(r, 10).value, wf.cell(r, 11).value
        if j is not None or k:
            escrever(ws, r, 10, (j or 0) + (k or 0), log)
        novo = horas.get(r)
        escrever(ws, r, 11, novo if novo else None, log)
        if novo:
            log.append(f"   L{r} {wf.cell(r, 1).value} {wf.cell(r, 2).value.strip()}: {novo:g} h")
    escrever(ws, 35, 14, n_ant + p_ant, log)
    escrever(ws, 2, 20, f"BMS_4600000002_GSV_BM_{num}_{rot}", log)
    escrever(ws, 2, 23, data_excel(C.DATA_BM), log)
    escrever(ws, 3, 21, f"{C.PERIODO_INI:%d/%m/%Y} a {C.PERIODO_FIM:%d/%m/%Y}", log)
    xl.app.CalculateFull()
    log.append(f"   N.º BM {ws.Range('T2').Value} | Data BM {ws.Range('W2').Text} | período {ws.Range('U3').Value}")
    log.append(f"   R$ anterior (N35) {n_ant:,.2f} + {p_ant:,.2f} = {n_ant + p_ant:,.2f} | no mês (P35) "
               f"{ws.Range('P35').Value:,.2f} | acumulado (R35) {ws.Range('R35').Value:,.2f}")
    wb.Save()
    wb.Close()


def fotos(xl, f, pasta_fotos, num, rot, log):
    """Cabeçalho do FOTOS + fotos do mês (subpastas por dia num temporário, para o rodízio do fotos_bm)."""
    wb = xl.abrir(f)
    aba = wb.Worksheets("FOTOS")
    cab = next((aba.Cells(r, c) for r in range(1, 5) for c in range(1, 33)
                if str(aba.Cells(r, c).Value or "").startswith("BMS_4600000002")), None)
    if cab is None:
        log.append("   !! FOTOS: cabeçalho 'BMS_4600000002…' não achado — N.º BM não alterado")
    else:
        cab.Value = f"BMS_4600000002_GSV_BM_{num}_{rot}"
    area = aba.Range(aba.PageSetup.PrintArea) if aba.PageSetup.PrintArea else aba.UsedRange
    lim_dir, lim_baixo = area.Left + area.Width, area.Top + area.Height
    vagas = sum(1 for s in aba.Shapes if s.Type == 13 and s.Height >= LOGO_ALTURA and s.Left < lim_dir and s.Top < lim_baixo)
    wb.Save()
    wb.Close()
    por_dia = defaultdict(list)
    for x in glob.glob(os.path.join(pasta_fotos, "*.jp*g")) + glob.glob(os.path.join(pasta_fotos, "*.png")):
        m = re.search(r"(\d{4}-\d{2}-\d{2})", os.path.basename(x))
        por_dia[m.group(1) if m else "~sem data"].append(x)
    if not por_dia:
        log.append("   !! sem fotos na pasta Fotos — FOTOS.xlsx ficou com as fotos antigas")
        return
    dias = dias_espalhados(sorted(por_dia), vagas)
    log.append(f"   FOTOS: {vagas} vagas; dias com foto {len(por_dia)} -> usados {', '.join(d[8:] + '/' + d[5:7] for d in dias)}")
    with tempfile.TemporaryDirectory() as tmp:
        for d in dias:
            os.makedirs(os.path.join(tmp, d))
            for x in por_dia[d]:
                shutil.copy2(x, os.path.join(tmp, d))
        log += preencher_fotos(f, tmp)


def main(numeros):
    horas = horas_caderno()
    log = [f"BOLETINS PREPARAÇÃO DE CARGAS — BM {C.BM_NUMERO} ({C.PERIODO_INI:%d/%m} a {C.PERIODO_FIM:%d/%m}), "
           f"Data BM {C.DATA_BM:%d/%m/%Y}, rolagem por equipe"]
    with Excel() as xl:
        for (eq, (sub, rot, rot_f)), num in zip(EQUIPES.items(), numeros):
            pasta = os.path.join(PASTA, sub)
            bol = [x for x in glob.glob(os.path.join(pasta, "Boletim*.xlsx")) if not os.path.basename(x).startswith("~$")]
            if len(bol) != 1:
                raise SystemExit(f"!! {pasta}: esperava 1 boletim, achei {bol}")
            bak = os.path.join(BACKUP, sub)
            os.makedirs(bak, exist_ok=True)
            for x in (bol[0], os.path.join(pasta, "FOTOS.xlsx")):
                if os.path.exists(x) and not os.path.exists(os.path.join(bak, os.path.basename(x))):
                    shutil.copy2(x, bak)
            log.append(f"\n=== {eq} — BM_{num} — {os.path.basename(bol[0])}")
            boletim(xl, bol[0], os.path.join(bak, os.path.basename(bol[0])), horas.get(eq, {}), num, rot, log)
            if os.path.exists(os.path.join(pasta, "FOTOS.xlsx")):
                fotos(xl, os.path.join(pasta, "FOTOS.xlsx"), os.path.join(pasta, "Fotos"), num, rot_f, log)
    sem = set(horas) - set(EQUIPES)
    if sem:
        log.append(f"!! equipes no Caderno sem boletim: {sem}")
    txt = "\n".join(log)
    open(os.path.join(BACKUP, "relatorio_boletins.txt"), "w", encoding="utf-8").write(txt)
    print(txt)


def teste():
    d = dias_espalhados(list(range(13)), 12)
    assert len(set(d)) == 12 and d[0] == 0 and d[-1] == 12, d
    assert dias_espalhados([1, 2], 5) == [1, 2]
    assert dias_espalhados(list(range(9)), 1) == [0]
    print("ok")


if __name__ == "__main__":
    if sys.argv[1:] == ["teste"]:
        teste()
    elif len(sys.argv) == 4 and all(re.fullmatch(r"\d{3}", a) for a in sys.argv[1:]):
        main(sys.argv[1:])
    else:
        raise SystemExit(__doc__)
