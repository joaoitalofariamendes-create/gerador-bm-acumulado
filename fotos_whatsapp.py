# -*- coding: utf-8 -*-
"""Separa as fotos de um grupo do WhatsApp (exportação .zip "com mídia") por PRT.

Para cada foto: quem enviou + dia -> colaborador (aba COLABORADORES) -> projeto em que ele foi
apontado nesse dia (aba APONTAMENTOS) -> pasta da PRT. Sem certeza (remetente desconhecido, sem
apontamento no dia, mais de um projeto no dia) -> pasta 'A REVISAR' com o motivo no relatório.
Uso: python fotos_whatsapp.py <export.zip> [dd/mm/aaaa início] [dd/mm/aaaa fim] [pasta de saída]
"""
import os
import re
import sys
import unicodedata
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime

import openpyxl

APONT = r"<DRIVE_MEDICAO>\Medições\Apontamentos\Apontamentos.xlsx"
SAIDA = r"<PASTA_TRABALHO>\Medicoes\fotos_whatsapp_teste"
# remetente no WhatsApp -> nome como no cadastro (COLABORADORES), quando o automático não resolve
APELIDOS = {}
RUIDO = {"GSV", "CORP", "CORPORATIVO", "ANDAIME", "ALMOX", "ENC", "ENCARREGADO"}
RDO_EQUIPE = {"ENCARREGADO A": "EQUIPE FIXA", "ENCARREGADO B": "CALDEIRARIA"}
MSG = re.compile(r"^\[(\d{2}/\d{2}/\d{4}), (\d{2}):(\d{2}):(\d{2})\] ([^:]+): (.*)$")


def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().upper()
    return " ".join(re.sub(r"[^A-Z ]", " ", s).split())


def pasta_projeto(projeto):
    m = re.search(r"PRT_?\s*(\d{4})", projeto)
    if m:
        return f"PRT {m.group(1)}"
    return re.sub(r'[\\/:*?"<>|]', "-", projeto.replace("CLIENTE - ", "")).strip()


def ler_apontamentos():
    wb = openpyxl.load_workbook(APONT, data_only=True, read_only=True)
    cad = [str(r[1]).strip() for r in wb["COLABORADORES"].iter_rows(min_row=2, values_only=True) if r[1]]
    dia = defaultdict(set)                      # (nome, data) -> {projeto}
    funcs = defaultdict(set)                    # nome -> {função apontada}
    for r in wb["APONTAMENTOS"].iter_rows(min_row=2, values_only=True):
        nome, dt, projeto, rdo = r[0], r[1], r[4], (r[7] if len(r) > 7 else None)
        if not (nome and projeto and isinstance(dt, datetime)):
            continue
        projeto = str(projeto).strip()
        if projeto == "CLIENTE - PREPARAÇÃO DE CARGAS" and rdo:
            autor = norm(str(rdo)).replace(" PDF", "").split()[-1]
            projeto += f" - {RDO_EQUIPE[autor]}" if autor in RDO_EQUIPE else ""
        dia[(str(nome).strip(), dt.date())].add(projeto)
        funcs[str(nome).strip()].add(norm(r[6] or ""))
    wb.close()
    return cad, dia, funcs


LIDER = ("ENCARREGADO", "SUPERVISOR", "TECNICO", "ENGENHEIRO")


def escolher(cands, funcs):
    """Entre vários, fica o único que é encarregado/supervisor/técnico/engenheiro (quem manda as fotos)."""
    lid = [c for c in cands if any(f.startswith(LIDER) for f in funcs.get(c, ()))]
    return lid[0] if len(lid) == 1 else None


def candidatos(remetente, cad):
    if remetente in APELIDOS:
        return [APELIDOS[remetente]]
    toks = [t for t in norm(remetente).split() if t not in RUIDO and len(t) > 1]
    if not toks:
        return []
    return [c for c in cad if all(t in norm(c).split() for t in toks)]


def main(zp, ini, fim, saida):
    cad, dia, funcs = ler_apontamentos()
    periodo = defaultdict(set)                  # nome -> {projeto} no período
    for (n, d), ps in dia.items():
        if ini <= d <= fim:
            periodo[n] |= ps
    z = zipfile.ZipFile(zp)
    txt = z.read("_chat.txt").decode("utf-8", "replace").replace("\u200e", "").replace("\u202f", " ")
    fotos = []
    for linha in txt.splitlines():
        m = MSG.match(linha)
        if m and "<anexado:" in m.group(6):
            arq = re.search(r"<anexado:\s*([^>]+)>", m.group(6)).group(1).strip()
            d = datetime.strptime(m.group(1), "%d/%m/%Y").date()
            if ini <= d <= fim and arq.lower().endswith((".jpg", ".jpeg", ".png")):
                fotos.append((d, f"{m.group(2)}.{m.group(3)}.{m.group(4)}", m.group(5).replace("~", "").strip(), arq))
    nomes_zip = {os.path.basename(n): n for n in z.namelist()}
    cont, revisar, quem, inferidas = Counter(), [], {}, []
    for d, hora, rem, arq in fotos:
        cands = candidatos(rem, cad)
        no_dia = [c for c in cands if (c, d) in dia]
        no_per = [c for c in cands if c in periodo]
        pessoa = (cands[0] if len(cands) == 1 else no_dia[0] if len(no_dia) == 1 else
                  no_per[0] if len(no_per) == 1 else escolher(no_dia or no_per, funcs))
        quem[rem] = pessoa or f"? ({len(cands)} candidatos)"
        motivo = None
        if pessoa is None:
            motivo = "remetente não identificado no cadastro" if not cands else f"remetente ambíguo: {', '.join(cands[:4])}"
        else:
            projs = dia.get((pessoa, d), set())
            if not projs and len(periodo.get(pessoa, ())) == 1:
                projs = periodo[pessoa]                 # único projeto dele no período
                dia[(pessoa, d)] = projs
                inferidas.append(f"{d:%d/%m} {rem} -> {pasta_projeto(next(iter(projs)))}")
            if not projs:
                motivo = f"{pessoa} sem apontamento em {d:%d/%m} e com {len(periodo.get(pessoa, ()))} projetos no período"
            elif len(projs) > 1:
                motivo = f"{pessoa} em mais de um projeto em {d:%d/%m}: {' | '.join(sorted(projs))}"
        destino = os.path.join(saida, "A REVISAR" if motivo else pasta_projeto(next(iter(dia[(pessoa, d)]))), f"{d:%d-%m}")
        os.makedirs(destino, exist_ok=True)
        nome = f"{d:%d-%m-%Y} {hora} - {rem} - {arq}"
        if arq not in nomes_zip:
            motivo = (motivo or "") + " | foto não está no zip"
        else:
            with open(os.path.join(destino, nome), "wb") as f:
                f.write(z.read(nomes_zip[arq]))
        cont["A REVISAR" if motivo else pasta_projeto(next(iter(dia[(pessoa, d)])))] += 1
        if motivo:
            revisar.append(f"{nome}: {motivo}")
    rel = [f"FOTOS {ini:%d/%m/%Y} a {fim:%d/%m/%Y} — {len(fotos)} fotos", "", "Por pasta:"]
    rel += [f"   {k}: {v}" for k, v in sorted(cont.items())]
    rel += ["", "Remetente -> colaborador:"] + [f"   {k} -> {v}" for k, v in sorted(quem.items())]
    rel += ["", f"INFERIDAS pelo único projeto da pessoa no período ({len(inferidas)}):"]
    rel += [f"   {k}: {v}" for k, v in sorted(Counter(inferidas).items())]
    rel += ["", f"A REVISAR ({len(revisar)}):"] + [f"   {x}" for x in revisar]
    with open(os.path.join(saida, "relatorio_fotos.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(rel))
    print("\n".join(rel[:40]))


if __name__ == "__main__":
    a = sys.argv[1:]
    ini = datetime.strptime(a[1], "%d/%m/%Y").date() if len(a) > 1 else date(2026, 9, 6)
    fim = datetime.strptime(a[2], "%d/%m/%Y").date() if len(a) > 2 else date(2026, 10, 5)
    main(a[0], ini, fim, a[3] if len(a) > 3 else SAIDA)
