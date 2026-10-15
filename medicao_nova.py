# -*- coding: utf-8 -*-
"""Medição-modelo para PRT que tem BM/Acumulado de meses anteriores mas nunca teve Medição
(a Medição só existe a partir do BM 5030 — ex.: 5001, 5026).

Copia a Medição do mês anterior que tem MAIS itens em comum com o Acumulado da PRT e:
  - cabeçalho: PRT (I6) e nome (E6) do Acumulado; ID Planview / NIQ / centro de custo /
    gestor (C6, D6, G6, H6) ficam em branco para preencher;
  - linhas de item que não são da PRT: ocultas e sem quantidade;
  - itens da PRT que o modelo não tem: linha CRIADA copiando uma linha irmã (as fórmulas
    PROCV da PQ puxam descrição, unidade e preço pelo código da coluna C).
Horas extras (1.5.x) não entram aqui: o aplicar_base cria as que tiverem valor no mês.
Grava em SAIDA\\_modelos\\<pasta da PRT>\\ (nunca no Drive).
"""
import glob
import os
import re
import shutil

import config_mes as C
from comum import (ADICIONAIS, Excel, acertar_totais_medicao, alinhar_cabecalho, arquivo, data_br, data_excel, eh_codigo, fonte,
                   ler_acumulado, ler_medicao, ordem, pasta_prt, procv_pq)

XL_DOWN = -4121


def medicao_modelo(prt):
    acu = fonte(prt, "acumulado")
    px = None if acu else C.PRT_NOVA[prt]          # PRT sem histórico: itens do xlsx da PRT
    pasta = os.path.basename(os.path.dirname(acu)) if acu else f"PRT {prt}"
    dst_dir = os.path.join(C.SAIDA, "_modelos", pasta)
    dst = os.path.join(dst_dir, f"Medicao_PRT_{prt}.xlsx")
    if os.path.exists(dst):
        return dst
    if acu:
        itens = {l["cod"] for l in ler_acumulado(acu)[1]
                 if l["und"] not in (None, "", 0) and l["cod"] not in ADICIONAIS}
    else:
        from escopo_pdf import ler_escopo_xlsx
        itens = {k for k in ler_escopo_xlsx(px) if "|" not in k}   # horas extras entram pelo aplicar_base
    # modelo: Medição do mês anterior com mais itens em comum
    melhor = None
    for p in glob.glob(os.path.join(C.FONTE_MEDICAO, "PRT*")):
        m = arquivo(p, "medicao")
        if not m:
            continue
        cods = {l["cod"] for l in ler_medicao(m) if l["und"] not in (None, "", 0)}
        n = len(cods & itens)
        if melhor is None or n > melhor[0]:
            melhor = (n, m)
    os.makedirs(dst_dir, exist_ok=True)
    shutil.copy2(melhor[1], dst)
    ident, nome = cabecalho(acu) if acu else cabecalho_prt(px)
    linhas = ler_medicao(dst)
    log = [f"PRT {prt}: Medição-modelo a partir de {os.path.basename(os.path.dirname(melhor[1]))}"
           f" ({melhor[0]} de {len(itens)} itens em comum)"]
    with Excel() as xl:
        wb = xl.abrir(dst)
        ws = wb.Worksheets("Formuário_PRT")
        alinhar_cabecalho(ws, log)               # modelo com a linha 6 deslocada (5008 do BM 5030)
        ws.Range("I6").Value = ident
        ws.Range("E6").Value = nome
        for cel in ("C6", "D6", "G6", "H6"):
            ws.Range(cel).Value = None
        if px:                                   # PRT nova: ID/NIQ/centro de custo/gestor vêm da PRT
            import openpyxl
            wp = openpyxl.load_workbook(px, data_only=True)["Formuário_PRT"]
            for cel in ("C6", "D6", "G6", "H6"):
                ws.Range(cel).Value = wp[cel].value
        ws.Range("J6").Value = data_excel(C.PERIODO_FIM)   # texto "05/10/2026" o Excel lia como 10/05 (americano)
        data_br(ws.Range("J6"))                            # modelo em m/d/aaaa mostraria 10/5/2026
        # linhas que não são da PRT: ocultas e zeradas (títulos ficam como estão)
        for l in linhas:
            if l["und"] in (None, "", 0):
                continue
            fora = l["cod"] not in itens if l["cod"] not in ADICIONAIS else l["mae"] not in itens
            if fora:
                if ws.Cells(l["r"], 10).Value not in (None, "") and not ws.Cells(l["r"], 10).HasFormula:
                    ws.Cells(l["r"], 10).Value = None
                ws.Rows(l["r"]).Hidden = True
        # itens da PRT que o modelo não tem: de baixo para cima
        tem = {l["cod"] for l in linhas}
        cods_linhas = [(l["r"], l["cod"]) for l in linhas if l["cod"] not in ADICIONAIS]
        faltam = sorted((c for c in itens if c not in tem), key=ordem, reverse=True)
        for cod in faltam:
            pai = cod.rsplit(".", 1)[0]
            irmas = [r for r, c in cods_linhas if c.rsplit(".", 1)[0] == pai and eh_codigo(c)]
            antes = [r for r, c in cods_linhas if c.split(".")[0] == cod.split(".")[0] and ordem(c) < ordem(cod)]
            if not (irmas or antes):
                log.append(f"   {cod}: sem linha irmã para copiar — não criei (lance à mão)")
                continue
            modelo = (irmas or antes)[-1]
            cod_de = dict(cods_linhas)
            menores = [r for r in irmas + antes if ordem(cod_de[r]) < ordem(cod)]
            pos = max(menores) + 1 if menores else min(irmas)
            he = {l["r"] for l in linhas if l["cod"] in ADICIONAIS}
            while pos in he:                     # depois das horas extras da função anterior
                pos += 1
            ws.Rows(modelo).Copy()
            ws.Rows(pos).Insert(XL_DOWN)
            xl.app.CutCopyMode = False
            ws.Cells(pos, 3).Value = "'" + cod          # texto, como as outras linhas (senão 6.3 vira número)
            ws.Cells(pos, 10).Value = None
            ws.Cells(pos, 12).Formula = f"=J{pos}*K{pos}"
            if not cod.startswith("6."):
                procv_pq(ws, pos)                # a linha copiada pode ter preço fixo de outro item (6.1)
            ws.Rows(pos).Hidden = True
            log.append(f"   CRIADA linha {pos} {cod} (cópia da linha {modelo})")
            # linhas abaixo desceram 1
            cods_linhas = [(r + 1 if r >= pos else r, c) for r, c in cods_linhas] + [(pos, cod)]
            linhas = [dict(l, r=l["r"] + 1) if l["r"] >= pos else l for l in linhas]
        # 6.x (fornecimentos sob demanda): o preço é o da PRT (Acumulado), não o da PQ / Medição-modelo
        if acu:
            precos = precos_6x(acu)
            for r in range(12, 400):
                cod = str(ws.Cells(r, 3).Value or "").strip()
                if cod in precos and not ws.Cells(r, 11).HasFormula and ws.Cells(r, 11).Value2 != precos[cod]:
                    log.append(f"   {cod} (linha {r}): preço {ws.Cells(r, 11).Value2} -> {precos[cod]} (Acumulado da PRT)")
                    ws.Cells(r, 11).Value = precos[cod]
        acertar_totais_medicao(ws, log)
        wb.Save()
        wb.Close()
    log.append("   PREENCHER no cabeçalho: ID Planview, NIQ, centro de custo, gestor")
    with open(os.path.join(dst_dir, "leia-me.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(log))
    print("\n".join(log))
    return dst


def cabecalho_prt(px):
    """('PRT-GSV-DOP-B-26-5019', 'TÍTULO DA PRT') a partir do xlsx da PRT (I6 e a linha depois de 'Título da PRT:')."""
    import openpyxl
    ws = openpyxl.load_workbook(px, data_only=True)["Formuário_PRT"]
    nome = ""
    for r in range(12, 40):
        if str(ws.cell(r, 3).value or "").strip().upper().startswith("TÍTULO DA PRT"):
            nome = " ".join(str(ws.cell(r + 1, 3).value or "").split())
            break
    return str(ws["I6"].value).strip(), nome


def precos_6x(acu):
    """{código 6.x: preço unitário (F)} do BM Acumulado da PRT."""
    import openpyxl
    from comum import aba_acumulado
    ws = aba_acumulado(openpyxl.load_workbook(acu, data_only=True))
    return {str(ws.cell(r, 2).value).strip(): ws.cell(r, 6).value for r in range(12, ws.max_row + 1)
            if str(ws.cell(r, 2).value or "").strip().startswith("6.") and isinstance(ws.cell(r, 6).value, (int, float))}


def cabecalho(acu):
    """('PRT-GSV-DOP-B-26-5001', 'NOME DA PRT') a partir do Acumulado (T2 e A6)."""
    import openpyxl
    from comum import aba_acumulado
    ws = aba_acumulado(openpyxl.load_workbook(acu, data_only=True))
    t2 = str(ws["T2"].value or "")
    ident = re.sub(r"^BM_?", "", t2).strip()          # T2 vem 'BM_PRT-...' ou 'BMPRT-...'
    a6 = " ".join(str(ws["A6"].value or "").split())
    nome = a6.split(" - ", 1)[1] if " - " in a6 else a6
    return ident, nome
