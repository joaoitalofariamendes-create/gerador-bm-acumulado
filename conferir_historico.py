# -*- coding: utf-8 -*-
"""Conciliação do histórico, aba por aba (inclusive HE e abas ocultas), BM do mês x BM do mês passado.

Para cada aba de item ligada a uma linha do BM Acumulado:
  - ACUMULADO BMS ANTERIORES (soma) = Acum. Anterior (J) do Acumulado do mês;
  - histórico do mês = histórico do mês passado + medido no mês passado (o que o Acumulado do mês passado confirma);
  - J = Acum. Atual (L) e N = R$ Acumulado (R) do Acumulado do mês passado;
  - linha do mês (BM atual) = No Mês (K) do Acumulado.
Uso: python conferir_historico.py [PRT...]  ->  SAIDA\\Conferencia_Historico_BMnnnn.xlsx (só leitura dos BMs)
"""
import os
import re
import sys
import warnings

import openpyxl
from openpyxl.styles import Font, PatternFill

import config_mes as C
from comum import ADICIONAIS, aba_acumulado, arquivo, fonte, ler_acumulado, ler_medicao, num, pasta_prt
from aplicar_bm import mapa_abas
from conferir_mes import ler_base, resumo_bm

warnings.filterwarnings("ignore")
TOL = 1e-6


class _Aba:                      # adaptador openpyxl -> interface usada pelo mapa_abas
    def __init__(self, ws):
        self.ws, self.Name = ws, ws.title

    def Range(self, a):
        return type("R", (), {"Value": self.ws[a].value})()


def acumulado(caminho):
    """{chave: (J, K, L, N R$ anterior, R R$ acumulado)}"""
    if not caminho:
        return {}
    ws = aba_acumulado(openpyxl.load_workbook(caminho, data_only=True))
    out = {}
    for l in ler_acumulado(caminho)[1]:
        j, ll = num(l["anterior"]) or 0, num(l["atual"]) or 0
        out[l["chave"]] = (j, ll - j, ll, num(ws.cell(l["r"], 14).value) or 0, num(ws.cell(l["r"], 18).value) or 0)
    return out


def conciliar(prt, base):
    pasta = pasta_prt(C.SAIDA, prt, C.PASTA_DA_PRT)
    bm, acu, med = arquivo(pasta, "bm"), arquivo(pasta, "acumulado"), arquivo(pasta, "medicao")
    bm_ant, acu_ant = fonte(prt, "bm"), fonte(prt, "acumulado")
    lm = ler_medicao(med)
    hh = {l["cod"]: str(l["desc"]) for l in lm if str(l["und"]).strip().lower() == "hh"}
    validos = {l["chave"] for l in lm if l["cod"] in ADICIONAIS} | {k for k in base if "|" in k}
    wb = openpyxl.load_workbook(bm, read_only=False, data_only=True)
    chave_de = {a.Name.strip(): k for a, k in mapa_abas(type("W", (), {"Worksheets": [_Aba(s) for s in wb.worksheets]})(),
                                                        hh, validos)}
    wb.close()
    novo, velho = resumo_bm(bm), (resumo_bm(bm_ant) if bm_ant else {})
    a31, a30 = acumulado(acu), acumulado(acu_ant)
    linhas = []
    for nome, (cod, vis, bmn, q, hist) in novo.items():
        k = chave_de.get(nome)
        if not k or k not in a31:
            continue
        j, kk, ll, n, _ = a31[k]
        _, k30, l30, _, r30 = a30.get(k, (0, 0, 0, 0, 0))
        h31 = sum(hist)
        mes = q if str(bmn).split(".")[0] == C.BM_NUMERO else 0
        obs = []
        if abs(h31 - j) > TOL:
            obs.append(f"histórico {h31:g} ≠ Acum. Anterior {j:g}")
        if abs(mes - kk) > TOL:
            obs.append(f"linha do mês {mes:g} ≠ No Mês {kk:g}")
        if abs(j - l30) > TOL:
            obs.append(f"Acum. Anterior {j:g} ≠ Acum. Atual do mês passado {l30:g}")
        if abs(n - r30) > 0.005:
            obs.append(f"R$ Anterior {n:,.2f} ≠ R$ Acumulado do mês passado {r30:,.2f}")
        h30, bm30, q30 = None, None, None
        if nome in velho:
            _, _, bm30, q30, hist30 = velho[nome]
            h30 = sum(hist30)
            rolou = h31 - h30
            if abs(rolou - (q30 or 0)) > TOL and abs(rolou) > TOL:
                obs.append(f"entrou no histórico {rolou:g}, mas a linha do mês do BM passado tinha {q30:g}")
            elif abs(rolou) <= TOL and (q30 or 0) and abs(k30) > TOL:
                obs.append(f"BM passado mediu {q30:g} (Acumulado confirma {k30:g}) e não foi para o histórico")
            elif abs(rolou) <= TOL and (q30 or 0):
                obs.append(f"linha do mês do BM passado ({bm30} = {q30:g}) não consta no Acumulado — fora do histórico (certo)")
        linhas.append([prt, nome, "visível" if vis else "oculta", k, h30, bm30, q30, h31, j, l30, mes, kk, ll, n, r30,
                       "OK" if not obs or all("(certo)" in o for o in obs) else "VER", "; ".join(obs)])
    return linhas


def main():
    so = [a for a in sys.argv[1:] if re.fullmatch(r"\d{4}", a)]
    base = ler_base()
    todas = []
    for prt in sorted(base):
        if (so and prt not in so) or not any(base[prt].values()):
            continue
        print(f"PRT {prt}...", flush=True)
        todas += conciliar(prt, base[prt])
    dst = os.path.join(C.SAIDA, f"Conferencia_Historico_BM{C.BM_NUMERO}.xlsx")
    if so and os.path.exists(dst):     # filtrada: as outras PRTs ficam como estavam na planilha completa
        velhas = [list(r) for r in openpyxl.load_workbook(dst)["Abas"].iter_rows(min_row=2, values_only=True)
                  if r[0] and str(r[0]) not in so]
        todas = sorted(velhas + todas, key=lambda l: str(l[0]))      # estável: ordem das abas dentro da PRT fica
    wb = openpyxl.Workbook()
    res = wb.active
    res.title = "Resumo"
    res.append(["PRT", "abas conferidas", "OK", "VER"])
    for prt in sorted({l[0] for l in todas}):
        ls = [l for l in todas if l[0] == prt]
        res.append([prt, len(ls), sum(l[15] == "OK" for l in ls), sum(l[15] == "VER" for l in ls)])
    ws = wb.create_sheet("Abas")
    ws.append(["PRT", "Aba", "Aba no book", "Item (Acumulado)", "Histórico BM passado", "BM da linha do mês (passado)",
               "Qtd linha do mês (passado)", "Histórico agora", "Acum. Anterior J", "Acum. Atual L do mês passado",
               "Linha do mês agora", "No Mês K", "Acum. Atual L", "R$ Anterior N", "R$ Acumulado R do mês passado",
               "Situação", "Observação"])
    for l in todas:
        ws.append(l)
        if l[15] == "VER":
            ws.cell(ws.max_row, 16).fill = PatternFill("solid", fgColor="FFE699")
    for s in (res, ws):
        for c in s[1]:
            c.font = Font(bold=True)
        s.auto_filter.ref = s.dimensions
        s.freeze_panes = "A2"
    for col, w in zip("ABCDEFGHIJKLMNOPQ", (7, 14, 9, 16, 11, 11, 11, 11, 11, 11, 11, 10, 11, 14, 14, 9, 90)):
        ws.column_dimensions[col].width = w
    wb.save(dst)
    ver = [l for l in todas if l[15] == "VER"]
    print(f"\n{len(todas)} abas conferidas, {len(ver)} para ver:")
    for l in ver:
        print(f"   {l[0]} '{l[1]}' ({l[3]}): {l[16]}")
    print(f"\nPlanilha: {dst}")


if __name__ == "__main__":
    main()
