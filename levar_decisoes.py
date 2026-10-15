# -*- coding: utf-8 -*-
"""Leva para a Base nova (Caderno atualizado) as decisões tomadas na Base revisada do MESMO BM (ex.: a da simulação).

Uso: python levar_decisoes.py <Base nova> <Base revisada>      (grava a Base nova; relatório na tela)
  - linha em que a revisada tinha VALOR ≠ Caderno (decisão: veículo = 1 mês, item zerado...): mesmo VALOR na nova;
  - linha acrescentada à mão na revisada (não veio do Caderno, ex.: mobilização 0,7): acrescentada na nova;
  - linha que só existe na nova com valor: fica amarela, para decidir (item novo no Caderno).
"""
import sys

import openpyxl
from openpyxl.styles import PatternFill

AMARELO, VERDE = PatternFill("solid", fgColor="FFF2CC"), PatternFill("solid", fgColor="E2EFDA")
num = lambda x: x if isinstance(x, (int, float)) else None


def chave(r):
    return (str(r[0]).strip(), f"{str(r[2]).strip()}|{str(r[1]).strip()}" if r[2] else str(r[1]).strip())


def main(nova, revisada):
    rev = {chave(r): r for r in openpyxl.load_workbook(revisada, data_only=True)["Base"]
           .iter_rows(min_row=4, max_col=9, values_only=True) if r[0] and r[1]}
    wb = openpyxl.load_workbook(nova)
    ws = wb["Base"]
    vistos, novos, levadas = set(), [], []
    for row in ws.iter_rows(min_row=4, max_col=9):
        if not (row[0].value and row[1].value):
            continue
        k = chave([c.value for c in row])
        vistos.add(k)
        r = rev.get(k)
        cad, val = num(row[6].value), num(row[7].value)
        if r is None:
            if val and str(row[8].value).startswith("CONFERIR"):
                novos.append(f"{k[0]} {k[1]} {str(row[3].value)[:45]} = {val:g} ({str(row[8].value)[10:70]})")
            continue
        r_cad, r_val = num(r[6]), num(r[7])
        decidiu = (r_cad or 0) != (r_val or 0) or (val is None and r_val is not None)   # inclui "0 = não mede"
        if decidiu and not str(r[8]).startswith("OK — regra"):   # decisão (não a regra da PRT)
            row[7].value = r_val
            row[8].value = f"OK — decisão da Base revisada: {r_val if r_val is not None else 'vazio'} (Caderno {cad})"
            for c in row:
                c.fill = VERDE
            levadas.append(f"{k[0]} {k[1]} {str(row[3].value)[:45]}: {val} -> {r_val}")
    # situação que o montar_base escreve sozinho: a linha veio do Caderno da simulação — se o Caderno real (ou uma regra
    # nova, ex.: HE do Supervisor da 5042 virou hora normal) não a traz mais, NÃO volta como "acrescentada à mão"
    auto = lambda s: s == "Caderno" or s.startswith(("OK — regra", "Caderno não traz", "CONFERIR: adicional sem linha",
                                                     "CONFERIR: item sem linha", "CONFERIR: linha está OCULTA",
                                                     "CONFERIR: medido no mês passado"))
    for k, r in rev.items():
        if k not in vistos and num(r[7]) and not auto(str(r[8] or "")):   # acrescentada à mão na revisada
            ws.append(list(r[:8]) + [f"CONFERIR: só existia na Base revisada (acrescentada à mão) — {r[8] or ''}"])
            for c in ws[ws.max_row]:
                c.fill = AMARELO
            levadas.append(f"{k[0]} {k[1]} {str(r[3])[:45]}: acrescentada = {r[7]} (confirmar)")
    wb.save(nova)
    print(f"{len(levadas)} decisão(ões) levada(s) da Base revisada:")
    print("\n".join(f"   {x}" for x in levadas))
    print(f"{len(novos)} linha(s) nova(s) com valor (não estavam na revisada) — DECIDIR:")
    print("\n".join(f"   {x}" for x in novos))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
