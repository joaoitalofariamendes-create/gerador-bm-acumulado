# -*- coding: utf-8 -*-
"""PASSO 2 — aplica a Base do Mês na Medição e no BM Acumulado de cada PRT.

Rode:  python aplicar_base.py [caminho da Base]
Gera cópias preenchidas em SAIDA\\<pasta da PRT>\\ — os arquivos do Drive não são tocados.

Medição (Formuário_PRT):  zera a coluna Quant. (J) dos itens e escreve os valores da Base.
                          Linha oculta que recebe valor é reexibida.
BM Acumulado:             Acum. Anterior (J) = Acum. Atual (L) do mês passado;
                          No Mês (K) = valor da Base; período (V3) atualizado.
                          Item do mês que não existe: linha CRIADA copiando uma linha de item
                          irmã (fórmulas e formatação), com os títulos-pai que faltarem;
                          dados da PQ (descrição, unidade, preço, CMS) e quantidade contratada
                          da aba Escopo da Base.
                          PRT na 1ª medição (sem Acumulado no mês passado): Acumulado NOVO a
                          partir do Acumulado de outra PRT, só com as linhas do escopo.
"""
import glob
import os
import re
import shutil
import sys

import openpyxl

import config_mes as C
from comum import (fonte, ADICIONAIS, Excel, aba_acumulado, acertar_totais_medicao, alinhar_cabecalho, arquivo, chave_cod,
                   com_funcao, data_br, data_excel, drive_ok,
                   eh_codigo, escrever, tirar_ajustes, padronizar_fontes, ler_acumulado, ler_escopo, ler_medicao,
                   mostrar_casas, norm,
                   num, ordem, pasta_prt, pq_da_medicao, pq_interna, procv_pq, prt_info)

XL_UP = -4162


def txt_cod(cod):
    """Código como o Excel deve guardar: '4' número, '4.2' texto (senão vira 4,2)."""
    return int(cod) if "." not in cod else "'" + cod


def ler_base(caminho):
    ws = openpyxl.load_workbook(caminho, data_only=True)["Base"]
    base, erros = {}, []
    for i, r in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
        prt, cod, mae, _, _, _, _, valor, sit = (list(r) + [None] * 9)[:9]
        if not prt or not cod:
            continue
        prt, cod = str(prt).strip(), str(cod).strip()
        chave = f"{str(mae).strip()}|{cod}" if mae else cod
        if valor in (None, ""):
            if str(sit or "").startswith("CONFERIR") and "OCULTA" not in str(sit):
                erros.append(f"linha {i}: PRT {prt} {chave} está amarela e sem VALOR — preencha (0 se não mede)")
            continue
        if num(valor) is None:
            erros.append(f"linha {i}: PRT {prt} {chave} VALOR '{valor}' não é número")
            continue
        base.setdefault(prt, {})[chave] = valor
    return base, erros


def copiar(src, prt):
    destino = os.path.join(C.SAIDA, os.path.basename(os.path.dirname(src)))
    os.makedirs(destino, exist_ok=True)
    # 'BM 29_ACUMULADO - PRT_5001.xlsx' na pasta do BM 5030: o nº de BM no nome é de outro mês
    dst = os.path.join(destino, re.sub(r"^BM\s*\d+\s*_", "BM_", os.path.basename(src)))
    shutil.copy2(src, dst)
    return dst


# ----------------------------------------------------------------- criação de linhas (COM)
XL_DOWN = -4121


def cod_str(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip() if v is not None else ""


def linhas_com(ws):
    """(linha do TOTAL 1, [(linha, código, unidade, tem_formula_G)]) lendo a aba pelo Excel."""
    vals = ws.Range("A1:D600").Value
    tot = next(i for i, row in enumerate(vals, 1)
               if i > 11 and (str(row[0] or "").strip().upper().startswith("TOTAL 1") or cod_str(row[1]) == "P0"))
    itens = []
    for r in range(12, tot):
        cod = cod_str(vals[r - 1][1])
        if eh_codigo(cod):
            itens.append((r, cod, vals[r - 1][3], ws.Cells(r, 7).HasFormula))
    return tot, itens


def e_item(t):
    return t[2] not in (None, 0, "") and t[3]


def copiar_linha(xl, ws, src, dest):
    ws.Rows(src).Copy()
    ws.Rows(dest).Insert(XL_DOWN)   # argumentos POSICIONAIS: por nome o pywin32 ignora
    xl.app.CutCopyMode = False
    ws.Rows(dest).Hidden = False


def preencher(ws, r, cod, pq, item, qtd=None, valor=None):
    desc, und, preco, cms = pq.get(cod, (None, None, None, None))
    ws.Cells(r, 2).Value = txt_cod(cod)
    ws.Cells(r, 3).Value = desc
    if item:
        ws.Cells(r, 1).Value = cms
        ws.Cells(r, 4).Value = und
        ws.Cells(r, 5).Value = qtd
        ws.Cells(r, 6).Value = preco
        ws.Cells(r, 10).Value = None
        ws.Cells(r, 11).Value = valor
        # linha nova: fórmulas padrão (a linha copiada pode ter valores fixos ou R$ anterior de outro item)
        for col, f in (("G", "=F{r}*E{r}"), ("L", "=J{r}+K{r}"), ("M", '=IF(D{r}="","",L{r}/E{r})'),
                       ("N", "=J{r}*F{r}"), ("O", "0"), ("P", '=IFERROR(ROUND(F{r}*K{r},2),"")'), ("Q", "0"),
                       ("R", "=P{r}+N{r}"), ("S", '=IFERROR(O{r}+Q{r},"")'), ("T", '=IFERROR(R{r}+S{r},"")'),
                       ("U", '=IFERROR(IF(R{r}=0,0,R{r}/G{r}),"")'), ("V", '=IFERROR(E{r}-L{r},"")'),
                       ("W", '=IFERROR(G{r}-R{r},"")')):
            ws.Range(f"{col}{r}").Formula = f.format(r=r)


def destino(itens, tot, pai, cod, mae=None):
    """Linha antes da qual o novo código entra, respeitando a ordem dentro do pai."""
    if mae:
        im = next(i for i, t in enumerate(itens) if t[1] == mae)
        j = im + 1
        while j < len(itens) and itens[j][1] in ADICIONAIS and ordem(itens[j][1]) < ordem(cod):
            j += 1
        return itens[j][0] if j < len(itens) else itens[im][0] + (j - im)
    if pai is None:
        return tot
    ip = next(i for i, t in enumerate(itens) if t[1] == pai)
    prof = cod.count(".")
    j = ip + 1
    while j < len(itens) and (itens[j][1].startswith(pai + ".") or itens[j][1] in ADICIONAIS):
        c = itens[j][1]
        if c.startswith(pai + ".") and c.count(".") == prof and ordem(c) > ordem(cod):
            return itens[j][0]
        j += 1
    return itens[j - 1][0] + 1


def modelos(itens, und=None, adicional=False):
    """(linha-modelo de título, linha-modelo de item) dentro da própria aba."""
    tit = next((t[0] for t in itens if not e_item(t) and t[1] not in ADICIONAIS), None)
    cands = [t for t in itens if e_item(t) and (t[1] in ADICIONAIS) == adicional]
    cands = cands or [t for t in itens if e_item(t)]
    mesma = [t for t in cands if str(t[2]).strip().lower() == str(und or "").strip().lower()]
    return tit, (mesma or cands)[0][0] if cands else None


def criar_no_acumulado(xl, ws, chave, valor, pq, qtd, log):
    cod, mae = chave_cod(chave)
    tot, itens = linhas_com(ws)
    cods = {t[1] for t in itens}
    if mae:
        if mae not in cods:
            log.append(f"   Acumulado: {chave} = {valor} — a função {mae} também não existe; não criei a hora extra")
            return False
        _, mod = modelos(itens, adicional=True)
        d = destino(itens, tot, None, cod, mae)
        copiar_linha(xl, ws, mod, d)
        preencher(ws, d, cod, pq, True, qtd, valor)
        # na PQ a HE é o % adicional (0,75 / 1,0); no Acumulado o preço é o da função × (1 + %)
        rm = next(t[0] for t in itens if t[1] == mae)
        rm = rm + 1 if rm >= d else rm   # a linha nova entrou antes/depois da função?
        pct, p_mae = ws.Cells(d, 6).Value2, ws.Cells(rm, 6).Value2   # Value2: moeda vem como número
        if isinstance(pct, (int, float)) and isinstance(p_mae, (int, float)):
            # HE (1.5.1/1.5.2) = hora cheia + adicional; noturno/periculosidade (1.5.3/1.5.4) = só o adicional
            ws.Cells(d, 6).Value = p_mae * (pct if cod in ("1.5.3", "1.5.4") else 1 + pct)
        else:
            log.append(f"   Acumulado: {chave} — não consegui calcular o preço da hora extra (preencha F{d}; "
                       f"% = {pct!r}, preço da função = {p_mae!r})")
        log.append(f"   Acumulado: CRIADA linha {d} {chave} (hora extra de {mae}) = {valor}")
        return True
    partes = cod.split(".")
    anc = [".".join(partes[:i]) for i in range(1, len(partes))]
    for i, a in enumerate(anc):
        if a in cods:
            continue
        tit, _ = modelos(itens)
        d = destino(itens, tot, anc[i - 1] if i else None, a)
        copiar_linha(xl, ws, tit, d)
        for c in range(1, 23):
            ws.Cells(d, c).Value = None
        preencher(ws, d, a, pq, False)
        log.append(f"   Acumulado: CRIADO título {a} na linha {d}")
        tot, itens = linhas_com(ws)
        cods = {t[1] for t in itens}
    _, mod = modelos(itens, pq.get(cod, (0, None))[1])
    d = destino(itens, tot, anc[-1] if anc else None, cod)
    copiar_linha(xl, ws, mod, d)
    preencher(ws, d, cod, pq, True, qtd, valor)
    falta = "" if qtd is not None else " — SEM quantidade contratada (preencha E)"
    log.append(f"   Acumulado: CRIADA linha {d} {cod} = {valor} (contratado {qtd}){falta}")
    return True


def corrigir_somas(ws):
    """Somatórios do TOTAL GERAL (linha 10) e do TOTAL 1 passam a cobrir da linha 12 à última linha de item."""
    tot, _ = linhas_com(ws)
    rx = re.compile(r"(\$?[A-Z]{1,2}\$?)(\d+):(\$?[A-Z]{1,2}\$?)(\d+)")
    for r in (10, tot):
        for c in range(1, 24):
            cel = ws.Cells(r, c)
            if not cel.HasFormula:
                continue
            f = cel.Formula
            nf = rx.sub(lambda m: f"{m.group(1)}12:{m.group(3)}{tot - 1}"
                        if 12 <= int(m.group(2)) < tot and int(m.group(4)) < tot else m.group(0), f)
            if nf != f:
                cel.Formula = nf


def limpar_reajuste(ws, log):
    """Não há reajuste no contrato: valor digitado em O / Q (Reajuste anterior / no mês) nas
    linhas de item vira 0. Fórmulas (S, totais) ficam. H / I (reajuste do contratado, colunas ocultas)
    das linhas de item ficam vazias — lixo ali ('vb', 2, '=G51*F51') entra no total que soma G:I."""
    tot, itens = linhas_com(ws)
    n = 0
    for c in "OQ":
        for cel in ws.Range(f"{c}12:{c}{tot - 1}"):
            v = cel.Value2
            if not cel.HasFormula and isinstance(v, (int, float)) and v:
                cel.Value = 0
                n += 1
    if n:
        log.append(f"   Acumulado: {n} valores nas colunas de Reajuste (O/Q) zerados — vinham preenchidos no mês passado")
    lixo = []
    for r, cod, und, _ in itens:
        for c in "HI" if und not in (None, 0, "") else "":
            cel = ws.Range(f"{c}{r}")
            if str(cel.Formula).strip() not in ("", "0"):
                lixo.append(f"{cod} {c}{r} ({cel.Formula!r})")
                cel.Value = None
    if lixo:
        log.append(f"   Acumulado: colunas ocultas de reajuste (H/I) limpas: {', '.join(lixo)}")


def formulas_acumulado(ws, log):
    """Linha de item com valor fixo em L (Acum. Atual), P, R (R$ Acumulado), S (reajuste total), T (total
    acumulado) ou U (avanço financeiro) volta à fórmula padrão; TOTAL do contratado (G) soma só a coluna G."""
    tot, itens = linhas_com(ws)
    trocas = []
    for r, cod, und, _ in itens:
        if und in (None, 0, ""):
            continue
        for col, f in (("L", "=J{r}+K{r}"), ("P", '=IFERROR(ROUND(F{r}*K{r},2),"")'), ("R", "=P{r}+N{r}"),
                       ("S", '=IFERROR(O{r}+Q{r},"")'), ("T", '=IFERROR(R{r}+S{r},"")'),
                       ("U", '=IFERROR(IF(R{r}=0,0,R{r}/G{r}),"")')):
            cel = ws.Range(f"{col}{r}")
            if not cel.HasFormula:
                trocas.append(f"{cod} {col}{r} (era {cel.Value2})")
                cel.Formula = f.format(r=r)
    for r in (10, tot):                               # =SUM(G12:I130) somava o reajuste oculto (H/I)
        g = ws.Range(f"G{r}")
        nf = re.sub(r"(\$?G\$?\d+):\$?I(\$?\d+)", r"\1:G\2", str(g.Formula))
        if g.HasFormula and nf != g.Formula:
            trocas.append(f"total G{r} (era {g.Formula})")
            g.Formula = nf
    if trocas:
        log.append(f"   Acumulado: fórmula padrão reposta em {len(trocas)} célula(s) com valor fixo: {', '.join(trocas)}")


def tirar_vinculos(ws, log):
    """Fórmula que lê OUTRO arquivo: abaixo do TOTAL é conferência solta (ex.: 5042 N134 contra uma proposta
    'DESATUALIZADO') e sai; nas linhas do BM fica e vai para o relatório."""
    tot, _ = linhas_com(ws)
    rng = ws.UsedRange
    r0, c0 = rng.Row, rng.Column
    for i, linha in enumerate(rng.Formula):
        for j, f in enumerate(linha):
            if not (isinstance(f, str) and f.startswith("=") and re.search(r"\[[^\]]+\.xls", f)):
                continue
            cel = ws.Cells(r0 + i, c0 + j)
            end = cel.Address.replace("$", "")
            if r0 + i > tot:
                log.append(f"   Acumulado: {end} lia outro arquivo (conferência solta, valor {cel.Value2}) — apagada: {f[:150]}")
                cel.Value = None
            else:
                log.append(f"   !! Acumulado: {end} lê outro arquivo — conferir: {f[:150]}")


def saldo_acumulado(ws, log):
    """Saldo: linha de item V = E-L (qtd) e W = G-R (R$), em fórmula; TOTAL 1 e linha 10: W = soma do saldo R$
    (o saldo em quantidade não tem total — soma unidades diferentes)."""
    tot, itens = linhas_com(ws)
    trocas = []
    for r, cod, und, _ in itens:
        if und in (None, 0, ""):
            continue
        m = ws.Range(f"M{r}")                        # avanço físico: sem contratado (E vazio/0) fica em branco
        if "IFERROR" not in str(m.Formula).upper():
            trocas.append(f"{cod} M{r} (avanço físico sem proteção: {m.Formula!r})")
            m.Formula = f'=IFERROR(L{r}/E{r},"")'
        for col, f in (("V", '=IFERROR(E{r}-L{r},"")'), ("W", '=IFERROR(G{r}-R{r},"")')):
            cel = ws.Range(f"{col}{r}")
            if not cel.HasFormula:
                trocas.append(f"{cod} {col}{r} (era {cel.Value2})")
                cel.Formula = f.format(r=r)
    soma = f"=SUM(W12:W{tot - 1})"
    for r in (10, tot):
        if str(ws.Range(f"W{r}").Formula).replace(" ", "") != soma:
            trocas.append(f"total W{r} (era {ws.Range(f'W{r}').Formula!r})")
            ws.Range(f"W{r}").Formula = soma
    if trocas:
        log.append(f"   Acumulado: saldo (V/W) reposto em {len(trocas)} célula(s): {', '.join(trocas)}")
    # saldo negativo (medido além do contratado) em vermelho e negrito: formatação condicional, vale também
    # para o que for editado à mão depois (regra do usuário, 13/10/2026)
    faixa = ws.Range(f"V12:W{tot}")
    faixa.FormatConditions.Delete()
    fc = faixa.FormatConditions.Add(1, 6, "=0")      # xlCellValue, xlLess
    fc.Font.Color = 255                              # vermelho (BGR)
    fc.Font.Bold = True
    negativos = [cod for r, cod, und, _ in itens if und not in (None, 0, "")
                 and any(isinstance(v, (int, float)) and v < -0.005 for v in (ws.Range(f"V{r}").Value, ws.Range(f"W{r}").Value))]
    if negativos:
        log.append(f"   !! Acumulado: saldo NEGATIVO (medido além do contratado) em {', '.join(negativos)}")


def cms_da_pq(ws, pq, log):
    """Código CMS (coluna A) de cada linha de item = o da PQ da Medição para aquele código — a cópia do mês
    passado trazia linhas sem CMS ou com o CMS de outro item (regra do usuário, 13/10/2026)."""
    _, itens = linhas_com(ws)
    trocas, sem = [], set()
    for r, cod, und, _ in itens:
        if und in (None, 0, ""):
            continue
        cms = (pq.get(cod) or (None,) * 4)[3]
        if cms in (None, "", "-"):
            sem.add(cod)
            continue
        atual = ws.Cells(r, 1).Value
        if str(atual or "").strip() != str(cms).strip():
            trocas.append(f"{cod} (era {atual!r})")
            ws.Cells(r, 1).Value = cms
    if trocas:
        log.append(f"   Acumulado: código CMS acertado pela PQ em {len(trocas)} linha(s): {', '.join(trocas)}")
    if sem:
        log.append(f"   !! Acumulado: PQ sem código CMS para {', '.join(sorted(sem))} — preencher à mão")


def reescrever_totais(ws):
    """PRT nova: o TOTAL do modelo pode somar células soltas (vira #REF! ao apagar linhas).
    TOTAL 1 e TOTAL GERAL passam a somar direto da linha 12 à última linha de item."""
    tot, _ = linhas_com(ws)
    u = tot - 1
    for c in "GNOPQS":
        ws.Range(f"{c}{tot}").Formula = f"=SUM({c}12:{c}{u})"
        ws.Range(f"{c}10").Formula = f"=SUM({c}12:{c}{u})"
    ws.Range(f"R{tot}").Formula = f"=N{tot}+P{tot}"
    ws.Range(f"T{tot}").Formula = f"=SUM(R{tot}:S{tot})"
    ws.Range(f"U{tot}").Formula = f'=IF(G{tot}>0,R{tot}/G{tot},"")'
    ws.Range("R10").Formula = f"=SUM(R12:R{u})"
    ws.Range("T10").Formula = f"=SUM(T12:T{u})"


def modelo_acumulado_nova():
    """Acumulado do mês passado (aba 'MODELO BM') com menos linhas — base para PRT nova."""
    drive_ok(C.FONTE_ACUMULADO)
    melhor = None
    for f in glob.glob(os.path.join(C.FONTE_ACUMULADO, "PRT*", "*.xls*")):
        if "ACUMULADO" not in os.path.basename(f).upper() or os.path.basename(f).startswith("~$"):
            continue
        aba, linhas = ler_acumulado(f)
        tem_item = any(l["und"] not in (None, 0, "") and l["cod"] not in ADICIONAIS for l in linhas)
        tem_tit = any(l["und"] in (None, 0, "") for l in linhas)
        if aba == "MODELO BM" and tem_item and tem_tit and (melhor is None or len(linhas) < melhor[0]):
            melhor = (len(linhas), f)
    return melhor[1]


def novo_acumulado(xl, prt, med, valores, esc, log):
    """BM Acumulado de PRT na 1ª medição: copia outro Acumulado e refaz só as linhas do escopo."""
    modelo = modelo_acumulado_nova()
    destino_dir = os.path.join(C.SAIDA, os.path.basename(os.path.dirname(med)))
    os.makedirs(destino_dir, exist_ok=True)
    dst = os.path.join(destino_dir, f"BM_ACUMULADO - PRT_{prt}.xlsx")
    shutil.copy2(modelo, dst)
    aba, _ = ler_acumulado(modelo)
    pq = pq_da_medicao(med)
    # preço da PRT aprovada (aba Escopo da Base) só para item que a PQ não precifica (ex.: 6.1 '-');
    # hora extra nunca: é % sobre o preço da função
    esc_preco = {}
    we = openpyxl.load_workbook(os.path.join(C.SAIDA, f"Base_do_Mes_BM{C.BM_NUMERO}.xlsx"), data_only=True)
    if "Escopo" in we.sheetnames:
        for p_, cod_, mae_, _, _, preco_, *_ in we["Escopo"].iter_rows(min_row=4, values_only=True):
            if str(p_).strip() == prt and num(preco_) is not None:
                esc_preco[f"{str(mae_).strip()}|{str(cod_).strip()}" if mae_ else str(cod_).strip()] = num(preco_)
    ident, nome = prt_info(med)

    # sequência de linhas: escopo (ordem da PRT) + o que foi medido e não está no escopo
    chaves = list(esc) + sorted((k for k in valores if k not in esc and valores[k]),
                                key=lambda k: ordem(chave_cod(k)[1] or chave_cod(k)[0]))
    seq, postos = [], set()
    for k in chaves:
        cod, mae = chave_cod(k)
        if mae:
            if mae not in postos:
                log.append(f"   Acumulado: {k} — função {mae} fora do escopo; hora extra não criada")
                continue
            i = max(i for i, s in enumerate(seq) if s[1] == mae or (s[2] == mae))
            seq.insert(i + 1, ("item", cod, mae, k))
            continue
        partes = cod.split(".")
        for a in (".".join(partes[:i]) for i in range(1, len(partes))):
            if a not in postos:
                seq.append(("tit", a, None, a))
                postos.add(a)
        seq.append(("item", cod, None, k))
        postos.add(cod)

    wb = xl.abrir(dst)
    ws = wb.Worksheets(aba)
    tot, itens = linhas_com(ws)
    tit, mod = modelos(itens)
    for r in sorted((t[0] for t in itens if t[0] not in (tit, mod)), reverse=True):
        ws.Rows(r).Delete()
    # os dois modelos ficam em 12 e 13; cada linha nova entra acima deles
    t_ofs, i_ofs = (0, 1) if tit < mod else (1, 0)
    pos = {}
    for n, (tipo, cod, mae, k) in enumerate(seq):
        dest = 12 + n
        copiar_linha(xl, ws, dest + (t_ofs if tipo == "tit" else i_ofs), dest)
        if tipo == "tit":
            for c in range(1, 23):
                ws.Cells(dest, c).Value = None
        preencher(ws, dest, cod, pq, tipo == "item", esc.get(k), valores.get(k) or None)
        if tipo != "item":
            continue
        if not mae:
            pos[cod] = dest
        if not mae and k in esc_preco and num(pq.get(cod, (0, 0, None))[2]) is None:
            ws.Cells(dest, 6).Value = esc_preco[k]       # ex.: 6.1 — a PQ não tem preço, vale o da PRT
        elif mae:          # hora extra = preço da função × (1 + %), como no criar_no_acumulado
            pct, p_mae = ws.Cells(dest, 6).Value2, ws.Cells(pos[mae], 6).Value2
            if isinstance(pct, (int, float)) and isinstance(p_mae, (int, float)):
                ws.Cells(dest, 6).Value = p_mae * (pct if cod in ("1.5.3", "1.5.4") else 1 + pct)
            else:
                log.append(f"   Acumulado: {k} — não consegui calcular o preço da hora extra (preencha F{dest})")
    ws.Rows(f"{12 + len(seq)}:{13 + len(seq)}").Delete()
    reescrever_totais(ws)
    limpar_reajuste(ws, log)
    formulas_acumulado(ws, log)
    tirar_vinculos(ws, log)
    cms_da_pq(ws, pq, log)
    saldo_acumulado(ws, log)
    padronizar_fontes(wb, "Acumulado", log)
    ws.Range("T2").Value = f"BM_{ident}"
    ws.Range("A6").Value = f"{ident} - {nome}"
    ws.Range("V3").Value = f"{C.PERIODO_INI:%d/%m/%Y} A {C.PERIODO_FIM:%d/%m/%Y}"
    data_do_bm(ws, log)
    casas_acumulado(ws, log)
    wb.Save()
    wb.Close()
    sem = [s[3] for s in seq if s[0] == "item" and esc.get(s[3]) is None]
    log.append(f"   Acumulado NOVO (1ª medição) a partir de {os.path.basename(modelo)}: "
               f"{sum(1 for s in seq if s[0] == 'item')} itens, {sum(1 for s in seq if s[0] == 'tit')} títulos")
    if sem:
        log.append(f"   Acumulado: sem quantidade contratada (coluna E vazia): {', '.join(sem)}")
    return dst


def linhas_medicao_com(ws, ini=12, fim=400):
    """[(linha, código, unidade)] da Formuário_PRT lidos pelo Excel (posições atuais)."""
    c = ws.Range(f"C{ini}:C{fim}").Value
    u = ws.Range(f"I{ini}:I{fim}").Value
    return [(ini + i, str(cc).strip(), uu) for i, ((cc,), (uu,)) in enumerate(zip(c, u)) if eh_codigo(cc)]


def aplicar_medicao(xl, src, valores, log, acu=None):
    linhas = ler_medicao(src)
    por_chave = {l["chave"]: l for l in linhas}
    dst = copiar(src, None)
    wb = xl.abrir(dst)
    ws = wb.Worksheets("Formuário_PRT")
    alinhar_cabecalho(ws, log)
    if data_br(ws.Range("J6")):
        log.append("   Medição: data do cabeçalho (J6) estava em formato americano (m/d/aaaa) — agora dd/mm/aaaa")
    limpas = escritas = 0
    for l in linhas:
        cel = ws.Cells(l["r"], 10)
        so_conta = cel.HasFormula and not re.search(r"[A-Za-z]", cel.Formula[1:])
        # só linha de item com unidade (col. I): na tabela de equipamentos (3.x) a col. J é
        # "Estimativa de uso" da PRT, não quantidade do mês — não zera
        if l["und"] not in (None, "", 0) and num(l["qtd"]) is not None and (not cel.HasFormula or so_conta):
            ws.Cells(l["r"], 10).Value = None
            limpas += 1
    # tabela 'Estimativa de uso' (3.x, unidade na coluna D e L = E×J×K) quando é usada para MEDIR (ex.: 5028):
    # entra no Custo Total da PRT, então a quantidade do mês passado não pode ficar — vale a da Base (ou 0)
    for r, cod, und in linhas_medicao_com(ws):
        f = str(ws.Cells(r, 12).Formula).replace("$", "")
        q = ws.Cells(r, 10).Value2
        if und in (None, "", 0) and f"E{r}" in f and f"J{r}" in f and isinstance(q, (int, float)) and q:
            novo = valores.get(cod) or 0
            if q != novo:
                ws.Cells(r, 10).Value = novo
                log.append(f"   Medição: linha {r} ({cod}, tabela 'Estimativa de uso' que entra no total) {q:g} -> {novo:g}")
    faltam = []
    for chave, v in valores.items():
        l = por_chave.get(chave)
        if l is None:
            if v:
                faltam.append((chave, v))
            continue
        if escrever(ws, l["r"], 10, v, log):
            escritas += 1
            if l["oculta"] and v:
                ws.Rows(l["r"]).Hidden = False
                log.append(f"   Medição: linha {l['r']} ({chave}) estava oculta — reexibi")
    # Medição que mostra só o que foi medido (ex.: 5043: toda linha de item zerada fica oculta, só títulos à
    # vista): a linha medida no mês passado e não neste também some, como no padrão dela
    cods = [l["cod"] for l in linhas]
    titulo = lambda c: any(o.startswith(c + ".") for o in cods)
    item = lambda l: l["und"] not in (None, "", 0) and not titulo(l["cod"])
    if not any(item(l) and not l["oculta"] and not num(l["qtd"]) for l in linhas):
        for l in linhas:
            if item(l) and not l["oculta"] and not num(valores.get(l["chave"])):
                ws.Rows(l["r"]).Hidden = True
                log.append(f"   Medição: linha {l['r']} ({l['chave']}) sem quantidade no mês — oculta (padrão desta Medição)")
    # item/hora extra sem linha: cria copiando linha irmã (itens primeiro, depois HE; posições
    # relidas no Excel a cada inserção)
    for chave, v in sorted(com_funcao(faltam, por_chave), key=lambda kv: ("|" in kv[0], ordem(chave_cod(kv[0])[0]))):
        cod, mae = chave_cod(chave)
        rows = linhas_medicao_com(ws)                 # [(linha, código, unidade)]
        if mae:
            mr = next((r for r, c, u in rows if c == mae), None)
            if mr is None:
                log.append(f"   Medição: {chave} = {v} — função {mae} sem linha (não inseri)")
                continue
            modelo = next((r for r, c, _ in rows if c == cod), None) or \
                next((r for r, c, _ in rows if c in ADICIONAIS), None)
            pos = mr + 1
            while any(r == pos and c in ADICIONAIS for r, c, _ in rows):
                pos += 1
            # HE (1.5.1/1.5.2) = hora cheia + %; adicional noturno/periculosidade (1.5.3/1.5.4) = só o %
            # (igual ao preço do Acumulado: função × (1 + %) ou função × %)
            formula = ("=IFERROR((K{m}*K{p})*J{p},0)" if cod in ("1.5.3", "1.5.4")
                       else "=IFERROR(((K{m}*K{p})+K{m})*(J{p}),0)")
        else:
            itens = [(r, c) for r, c, u in rows if c not in ADICIONAIS and u not in (None, "", 0)]
            pai = cod.rsplit(".", 1)[0]
            irmas = [r for r, c in itens if c.rsplit(".", 1)[0] == pai]
            grupo = [r for r, c in itens if c.split(".")[0] == cod.split(".")[0]]
            cands = irmas or grupo or [r for r, _ in itens]    # sem irmã: qualquer linha de item
            modelo = cands[-1] if cands else None
            de = dict(itens)
            menores = [r for r in cands if ordem(de[r]) < ordem(cod)]
            pos = (max(menores) + 1) if menores else (min(cands) if cands else None)
            while pos and any(r == pos and c in ADICIONAIS for r, c, _ in rows):
                pos += 1
            formula, mr = "=J{p}*K{p}", None
        if modelo is None or pos is None:
            log.append(f"   Medição: {chave} = {v} — nenhuma linha irmã para copiar (não inseri)")
            continue
        ws.Rows(modelo).Copy()
        ws.Rows(pos).Insert(XL_DOWN)
        xl.app.CutCopyMode = False
        ws.Rows(pos).Hidden = False
        ws.Cells(pos, 3).Value = txt_cod(cod)
        ws.Cells(pos, 10).Value = v
        ws.Cells(pos, 12).Formula = formula.format(m=mr, p=pos)
        if not mae and not cod.startswith("6."):
            procv_pq(ws, pos)          # a linha copiada pode ser de outro bloco (ex.: 6.1 com preço fixo)
        elif cod.startswith("6."):
            # 6.x: o PROCV da linha copiada dá '-' (na PQ o 6.x não tem preço) e o total vira #VALOR! (5007, BM 5031);
            # o preço da PRT (F do Acumulado) entra logo abaixo, em precos_6x
            ws.Cells(pos, 11).Value = None
        escritas += 1
        log.append(f"   Medição: CRIADA linha {pos} {chave} = {v} (cópia da linha {modelo})")
    for r, cod, _ in linhas_medicao_com(ws):
        v = ws.Cells(r, 3).Value
        if isinstance(v, float) and not v.is_integer():       # 6.3 gravado como número: texto, como as outras
            ws.Cells(r, 3).Value = txt_cod(cod)                # (título 1, 4, 5… fica número)
    if acu:   # 6.x (fornecimento sob demanda): preço da PRT (F do Acumulado), não o da PQ/modelo (o da PQ é o de outra PRT)
        from medicao_nova import precos_6x
        precos = precos_6x(acu)
        for r, cod, _ in linhas_medicao_com(ws):
            k = ws.Cells(r, 11)
            if cod in precos and not k.HasFormula and k.Value2 != precos[cod]:
                log.append(f"   Medição: {cod} (linha {r}) preço {k.Value2} -> {precos[cod]} (Acumulado da PRT)")
                k.Value = precos[cod]
    acertar_totais_medicao(ws, log)
    pq_interna(wb, ws, log)
    tirar_ajustes(ws, "Medição", log)
    padronizar_fontes(wb, "Medição", log)
    wb.Save()
    wb.Close()
    log.append(f"   Medição: {escritas} valores escritos ({limpas} quantidades do modelo zeradas antes)")
    return dst


def acertar_total_anterior(ws, r_tot_ant, soma_itens_ant, form_ant, log):
    """TOTAL 1: R$ Anterior (N) = SOMA PURA das linhas (cada linha já traz o R$ Acumulado dela do mês
    passado). Nunca entra valor digitado para forçar o total: se não bater com o TOTAL do mês passado,
    vai para o relatório com a fórmula de lá, para corrigir na origem."""
    tot, _ = linhas_com(ws)
    antes = ws.Range(f"N{tot}").Formula
    nova = f"=SUM(N12:N{tot - 1})"
    ws.Range(f"N{tot}").Formula = nova
    if str(antes).replace(" ", "") != nova:
        log.append(f"   Acumulado: TOTAL N{tot} era '{antes}' — agora '{nova}' (soma pura das linhas)")
    if r_tot_ant is not None and abs(r_tot_ant - soma_itens_ant) > 1e-6:
        log.append(f"   !! Acumulado: R$ Anterior (soma das linhas) {soma_itens_ant:,.2f} ≠ R$ Acumulado do mês "
                   f"passado {r_tot_ant:,.2f} (diferença {r_tot_ant - soma_itens_ant:+,.2f}); TOTAL de lá era "
                   f"'{form_ant}' — conferir a origem")


def aplicar_acumulado(xl, src, valores, log, pq=None, esc=None):
    aba, linhas = ler_acumulado(src)
    por_chave = {l["chave"]: l for l in linhas}
    # R$ acumulado (coluna R) do mês passado, exato, item a item e no TOTAL 1
    wv = openpyxl.load_workbook(src, data_only=True)[aba]
    r_ant = {l["r"]: num(wv.cell(l["r"], 18).value) for l in linhas}
    dst = copiar(src, None)
    wb = xl.abrir(dst)
    ws = wb.Worksheets(aba)
    tot, _ = linhas_com(ws)
    r_tot_ant = num(wv.cell(tot, 18).value)
    form_ant = openpyxl.load_workbook(src)[aba].cell(tot, 18).value
    rolados, soma_ant = 0, 0
    for l in linhas:
        if num(l["atual"]) is None:
            continue
        if ws.Cells(l["r"], 12).HasFormula is False and num(l["qtd"]):
            log.append(f"   Acumulado: linha {l['r']} ({l['chave']}) Acum. Atual é valor fixo, não fórmula — conferir")
        if escrever(ws, l["r"], 10, l["atual"], log):
            rolados += 1
        if l["und"] not in (None, "") and r_ant.get(l["r"]) is not None:
            ws.Cells(l["r"], 14).Value = r_ant[l["r"]]      # R$ Anterior = R$ Acumulado do mês passado
            soma_ant += r_ant[l["r"]]
        k = ws.Cells(l["r"], 11)
        if not k.HasFormula and k.Value is not None:
            k.Value = None
    escritas, faltam = 0, []
    for chave, v in valores.items():
        l = por_chave.get(chave)
        if l is None:
            if v:
                faltam.append((chave, v))
            continue
        if escrever(ws, l["r"], 11, v, log):
            escritas += 1
    # itens novos: filhos antes das horas extras, na ordem dos códigos
    com_funcao(faltam, por_chave)
    faltam.sort(key=lambda kv: ("|" in kv[0], ordem(chave_cod(kv[0])[0])))
    criadas = sum(criar_no_acumulado(xl, ws, k, v, pq or {}, (esc or {}).get(k), log) for k, v in faltam)
    if criadas:
        corrigir_somas(ws)
    limpar_reajuste(ws, log)
    formulas_acumulado(ws, log)
    tirar_ajustes(ws, "Acumulado", log)
    tirar_vinculos(ws, log)
    padronizar_fontes(wb, "Acumulado", log)
    acertar_total_anterior(ws, r_tot_ant, soma_ant, form_ant, log)
    if pq:
        cms_da_pq(ws, pq, log)
    saldo_acumulado(ws, log)
    v3 = ws.Range("V3")
    if not v3.HasFormula:
        v3.Value = f"{C.PERIODO_INI:%d/%m/%Y} A {C.PERIODO_FIM:%d/%m/%Y}"
    data_do_bm(ws, log)
    casas_acumulado(ws, log)
    wb.Save()
    wb.Close()
    log.append(f"   Acumulado [{aba}]: {rolados} acumulados rolados, {escritas} valores do mês escritos")
    return dst


def casas_acumulado(ws, log):
    """Quantidades (E contratado, J/K/L) mostram todas as casas do valor — 0,38 não pode aparecer 0,4."""
    ws.Calculate()
    tot, _ = linhas_com(ws)
    feito = mostrar_casas(ws, f"A12:L{tot - 1}", (5, 10, 11, 12))
    if feito:
        log.append(f"   Acumulado: formato mostrava menos casas que o valor — {', '.join(feito)}")


def data_do_bm(ws, log):
    """'Data BM:' do cabeçalho do Acumulado = data de postagem do BM (config_mes.DATA_BM; sem ela, fica a do modelo)."""
    d = getattr(C, "DATA_BM", None)
    if not d:
        log.append("   Acumulado: config_mes.DATA_BM vazia — 'Data BM' ficou a do mês passado")
        return
    for c in range(1, 30):
        if norm(str(ws.Cells(2, c).Value or "")).startswith("DATA BM"):
            ws.Cells(2, c + 1).Value = data_excel(d)
            return
    log.append("   Acumulado: não achei 'Data BM:' na linha 2 — data do BM não atualizada")


def main():
    args = [a for a in sys.argv[1:] if not re.fullmatch(r"\d{4}", a)]
    so = [a for a in sys.argv[1:] if re.fullmatch(r"\d{4}", a)]
    caminho = args[0] if args else os.path.join(C.SAIDA, f"Base_do_Mes_BM{C.BM_NUMERO}.xlsx")
    base, erros = ler_base(caminho)
    if erros:
        print("A Base tem pendências — nada foi gerado:")
        for e in erros:
            print("  -", e)
        sys.exit(1)

    escopo = ler_escopo(caminho)
    log = [f"Base: {caminho}"]
    with Excel() as xl:
        for prt in sorted(so or base):
            print(f"PRT {prt}...", flush=True)
            log.append(f"\nPRT {prt}")
            if not any(base.get(prt, {}).values()):     # sem medição no mês: nada vai para a pasta
                log.append("   nenhuma quantidade no mês — Medição/Acumulado não gerados")
                continue
            med = fonte(prt, "medicao")
            acu = fonte(prt, "acumulado")
            esc = escopo.get(prt, {})
            if med:
                aplicar_medicao(xl, med, base[prt], log, acu)
            else:
                log.append("   Medição: sem modelo — não gerada")
            if acu:
                aplicar_acumulado(xl, acu, base[prt], log, pq_da_medicao(med) if med else {}, esc)
            elif med:
                novo_acumulado(xl, prt, med, base[prt], esc, log)
            else:
                log.append("   Acumulado: sem Acumulado anterior nem Medição — não gerado")
            print("\n".join(log[-1:]))
    rel = os.path.join(C.SAIDA, "relatorio_aplicacao.txt")
    with open(rel, "w", encoding="utf-8") as f:
        f.write("\n".join(log))
    print("\n".join(log))
    print(f"\nRelatório: {rel}")


if __name__ == "__main__":
    main()
