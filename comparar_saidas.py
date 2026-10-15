# -*- coding: utf-8 -*-
"""Compara duas saídas do mesmo BM célula a célula (ex.: a revisada x o gerador rodado do zero num rascunho).

Uso: python comparar_saidas.py <pasta A> <pasta B> [trecho do caminho...]  -> <pasta B>\\comparar.txt
Compara abas (ordem/visibilidade), valores e fórmulas, formato numérico, linhas ocultas e área de impressão.
Diferença = acerto feito à mão que o gerador não reproduz (ou defeito que ficou numa das pastas). Só leitura.
"""
import glob
import os
import sys
from collections import defaultdict

import openpyxl

LIM = 12


def v(x):
    x = getattr(x, "text", x)                   # ArrayFormula
    return None if x == "" else x


def igual(a, b):
    a, b = v(a), v(b)
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool):
        return abs(a - b) <= 1e-9 * max(1, abs(a), abs(b))
    if isinstance(a, str) and isinstance(b, str):
        return a.strip() == b.strip()
    return a == b


def comparar(fa, fb):
    wa, wb = openpyxl.load_workbook(fa), openpyxl.load_workbook(fb)
    out = []
    na = [(s.title, s.sheet_state) for s in wa.worksheets]
    nb = [(s.title, s.sheet_state) for s in wb.worksheets]
    if na != nb:
        sa, sb = dict(na), dict(nb)
        estado = [f"{n}: {sa[n]}->{sb[n]}" for n in sa if n in sb and sa[n] != sb[n]]
        ordem = [n for n, _ in na if n in sb] != [n for n, _ in nb if n in sa]
        out.append(f"   abas — só em A: {[n for n, _ in na if n not in sb]} | só em B: {[n for n, _ in nb if n not in sa]}"
                   f" | visibilidade: {estado}{' | ORDEM diferente' if ordem else ''}")
    for sa in wa.worksheets:
        if sa.title not in wb.sheetnames:
            continue
        sb = wb[sa.title]
        dif, fmt = [], defaultdict(int)
        mr, mc = max(sa.max_row, sb.max_row), max(sa.max_column, sb.max_column)
        for r in range(1, mr + 1):
            for c in range(1, mc + 1):
                ca, cb = sa.cell(r, c), sb.cell(r, c)
                if not igual(ca.value, cb.value):
                    dif.append(f"{ca.coordinate}: {str(v(ca.value))[:60]!r} | B {str(v(cb.value))[:60]!r}")
                elif ca.value is not None and ca.number_format != cb.number_format:
                    fmt[f"{ca.number_format} -> B {cb.number_format}"] += 1
        oc = [r for r in range(1, mr + 1) if bool(sa.row_dimensions[r].hidden) != bool(sb.row_dimensions[r].hidden)]
        if sa.print_area != sb.print_area:
            dif.append(f"área de impressão {sa.print_area} | B {sb.print_area}")
        if dif or fmt or oc:
            out.append(f"   [{sa.title}] {len(dif)} célula(s) diferente(s)"
                       + (f", {len(oc)} linha(s) oculta(s) diferente(s): {oc[:15]}{' …' if len(oc) > 15 else ''}" if oc else ""))
            out += [f"      {d}" for d in dif[:LIM]] + ([f"      … +{len(dif) - LIM}"] if len(dif) > LIM else [])
            out += [f"      formato ({n}x): {k}" for k, n in list(fmt.items())[:4]]
    return out


def main():
    a, b, *so = sys.argv[1:]
    rel = []
    for f in sorted(glob.glob(os.path.join(a, "**", "*.xlsx"), recursive=True)):
        r = os.path.relpath(f, a)
        if os.path.basename(f).startswith("~$") or (so and not any(p in r for p in so)):
            continue
        fb = os.path.join(b, r)
        if not os.path.exists(fb):
            rel.append(f"## {r}: só em A")
            continue
        dif = comparar(f, fb)
        rel.append(f"## {r}: {'IGUAL' if not dif else ''}")
        rel += dif
        print(rel[-1 - len(dif)], flush=True)
    for f in glob.glob(os.path.join(b, "**", "*.xlsx"), recursive=True):
        r = os.path.relpath(f, b)
        if not os.path.exists(os.path.join(a, r)) and (not so or any(p in r for p in so)):
            rel.append(f"## {r}: só em B")
    open(os.path.join(b, "comparar.txt"), "w", encoding="utf-8").write("\n".join(rel))
    print(f"\nRelatório: {os.path.join(b, 'comparar.txt')}")


if __name__ == "__main__":
    main()
