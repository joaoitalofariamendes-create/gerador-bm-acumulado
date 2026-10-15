# -*- coding: utf-8 -*-
"""Proposta (PDF/xlsx da PRT) x BM Acumulado (E quant., F preço) x BM de abas (PREVISTO).

Só leitura. Relatório: SAIDA/conferencia_previsto.txt. Com --corrigir, grava no BM de abas o PREVISTO
(célula abaixo do rótulo 'Previsto') = E do BM Acumulado (valor em célula que já existe).
Uso: python conferir_previsto.py [PRT...] [--corrigir]
"""
import glob
import os
import re
import sys
import warnings

import openpyxl

import config_mes as C
from aplicar_bm import mapa_abas, por_previsto
from comum import (ADICIONAIS, Excel, arquivo, aba_acumulado, casas_valor, eh_codigo, escrever, ler_acumulado,
                   ler_medicao, mostrar_casas, num)
from escopo_pdf import ler_escopo_pdf, prt_nova

warnings.filterwarnings("ignore")
TOL_Q, TOL_P = 1e-6, 0.005   # preço: a proposta em PDF mostra 2 casas


class _Aba:                      # adaptador openpyxl -> interface COM usada pelo mapa_abas
    def __init__(self, ws):
        self.ws, self.Name = ws, ws.title

    def Range(self, a):
        return type("R", (), {"Value": self.ws[a].value})()


def celula_previsto(ws):
    """Endereço do valor do PREVISTO: a célula abaixo do rótulo 'Previsto' (coluna U). O layout muda entre
    abas (U19/U20/U21; nas HE do modelo 'Executado Anterior' não há Previsto) — U20 fixo gravava em rótulo/fórmula."""
    for r in range(15, 23):
        if str(ws.cell(r, 21).value or "").strip().upper() == "PREVISTO":
            return f"U{r + 1}"
    return None


def esconde_casas(cel):
    """O formato mostra menos casas que o valor (1913,8332 aparecendo 1.913,83)."""
    s = re.sub(r'"[^"]*"|\\.|_.|\*.', "", str(cel.number_format or "General").split(";")[0])
    if not isinstance(cel.value, (int, float)) or s == "General" or "%" in s or not re.search(r"[0#]", s):
        return False
    m = re.search(r"\.([0#]+)", s)
    return casas_valor(cel.value) > (len(m.group(1)) if m else 0)


def proposta(pasta, prt):
    px = prt_nova(prt)
    if px:
        return px, ler_escopo_pdf(px)
    pdfs = [f for f in glob.glob(os.path.join(pasta, "*.pdf")) if "PRT" in os.path.basename(f).upper()]
    return (pdfs[0], ler_escopo_pdf(pdfs[0])) if pdfs else (None, {})


def conferir(pasta, prt, corrigir, log):
    acu, bm = arquivo(pasta, "acumulado"), arquivo(pasta, "bm")
    fonte, prop = proposta(pasta, prt)
    aba, linhas = ler_acumulado(acu)
    wv = openpyxl.load_workbook(acu, data_only=True)[aba]
    E = {l["chave"]: (num(wv.cell(l["r"], 5).value), num(wv.cell(l["r"], 6).value), l)
         for l in linhas if l["und"] not in (None, "", 0)}
    err = []
    # 1) proposta x Acumulado
    if not prop:
        err.append(f"sem proposta legível ({os.path.basename(fonte) if fonte else 'nenhum PDF na pasta'}) — só Acumulado x abas")
    for k, (und, q, p) in prop.items():
        if k not in E:
            if q:
                err.append(f"{k}: está na proposta ({q:g} {und}) e não tem linha no BM Acumulado")
            continue
        e, f, _ = E[k]
        if e is None or abs(e - q) > TOL_Q:
            err.append(f"{k}: quantidade proposta {q:g} ≠ Acumulado E {e}")
        if "|" not in k and p is not None and f is not None and abs(f - p) > TOL_P:
            err.append(f"{k}: preço proposta {p:g} ≠ Acumulado F {f:g}")
    for k, (e, f, l) in E.items():
        if prop and "|" not in k and k not in prop and l["cod"] not in ADICIONAIS and e:
            err.append(f"{k}: no BM Acumulado (E {e:g}) mas não na proposta")
    # 2) Acumulado x PREVISTO das abas
    fix = []
    if bm:
        wb = openpyxl.load_workbook(bm, data_only=True)
        desc_hh = {l["cod"]: str(l["desc"]) for l in linhas if str(l["und"]).lower() == "hh"}
        # aba de HE -> função: mesmo critério do aplicar_bm (só função com essa HE na Medição do mês) — sem ele,
        # '1.5.1 ME' (Mecânico Montador) ia para o Ajudante de Mecânica e o Previsto saía o da função errada
        med = arquivo(pasta, "medicao")
        validos = {l["chave"] for l in ler_medicao(med) if l["cod"] in ADICIONAIS} if med else set()
        for a, k in mapa_abas(type("W", (), {"Worksheets": [_Aba(s) for s in wb.worksheets]})(), desc_hh, validos):
            vis = wb[a.Name].sheet_state == "visible"
            if k is None or k not in E or (E[k][0] is None and not vis):
                continue
            pos = celula_previsto(a.ws)
            if pos is None:
                continue
            u20, e = num(a.ws[pos].value), E[k][0]
            if not e and vis:    # sem quantidade contratada: Previsto em branco (usuário, 03/10 — vem com a proposta revisada)
                if a.ws[pos].value not in (None, ""):
                    err.append(f"aba '{a.Name}' ({k}): PREVISTO {pos} {a.ws[pos].value} sem quantidade contratada "
                               "no Acumulado — fica em branco")
                    fix.append((a.Name, pos, None))
            elif u20 is None or abs(u20 - e) > TOL_Q:            # aba oculta com contratado 0: Previsto 0
                ok_prop = k in prop and abs(prop[k][1] - e) <= TOL_Q
                err.append(f"aba '{a.Name}' ({k}): PREVISTO {pos} {a.ws[pos].value} ≠ Acumulado E {e:g}"
                           + (" — proposta confirma o Acumulado" if ok_prop else " — proposta não confirma"))
                fix.append((a.Name, pos, e))       # o BM Acumulado é a referência do contratado
            elif esconde_casas(a.ws[pos]):
                err.append(f"aba '{a.Name}' ({k}): PREVISTO {pos} {e:g} aparece com menos casas (formato)")
                fix.append((a.Name, pos, e))
        wb.close()
    if corrigir and fix:
        with Excel() as xl:
            w = xl.abrir(bm)
            for nome, pos, e in fix:
                u = w.Worksheets(nome).Range(pos)
                antes = u.Formula if u.HasFormula else u.Value
                if e is None:
                    por_previsto(w.Worksheets(nome), None)          # em branco; avanço/saldo sem #DIV/0!
                    log.append(f"   CORRIGIDO aba '{nome}': PREVISTO {pos} {antes!r} -> em branco (sem contrato)")
                    continue
                u.Value = e
                mostrar_casas(w.Worksheets(nome), pos, (21,))      # 1913,8332 não pode aparecer 1.913,83
                log.append(f"   CORRIGIDO aba '{nome}': PREVISTO {pos} {antes!r} -> {e:g}")
            w.Save()
            w.Close()
    return fonte, err


def main():
    so = [a for a in sys.argv[1:] if re.fullmatch(r"\d{4}", a)]
    corrigir = "--corrigir" in sys.argv
    rel = [f"PROPOSTA x BM ACUMULADO x PREVISTO (BM de abas) — {C.SAIDA}"]
    for pasta in sorted(glob.glob(os.path.join(C.SAIDA, "PRT*"))):
        prt = re.search(r"(\d{4})", os.path.basename(pasta)).group(1)
        if so and prt not in so:
            continue
        log = []
        try:
            fonte, err = conferir(pasta, prt, corrigir, log)
        except Exception as e:
            fonte, err = None, [f"não consegui conferir: {e}"]
        rel.append(f"\nPRT {prt} ({os.path.basename(fonte) if fonte else '-'}): " + ("OK" if not err else f"{len(err)} ponto(s)"))
        rel += [f"   {x}" for x in err] + log
    with open(os.path.join(C.SAIDA, "conferencia_previsto.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(rel))
    print("\n".join(rel))


if __name__ == "__main__":
    main()
